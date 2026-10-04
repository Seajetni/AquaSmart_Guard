import docx
import os

docx_path = r"C:\Users\seaza\Downloads\AquaSmart_Guard.docx"
output_txt = r"D:\aquarium-dashboard\docx_content.txt"

doc = docx.Document(docx_path)

with open(output_txt, "w", encoding="utf-8") as f:
    f.write(f"Number of paragraphs: {len(doc.paragraphs)}\n")
    f.write(f"Number of tables: {len(doc.tables)}\n\n")
    f.write("=== PARAGRAPHS ===\n")
    for i, p in enumerate(doc.paragraphs):
        style_name = p.style.name if p.style else "NoStyle"
        text = p.text.strip()
        if text:
            f.write(f"[{i}] ({style_name}): {p.text}\n")
    
    f.write("\n=== TABLES ===\n")
    for t_idx, table in enumerate(doc.tables):
        f.write(f"\n--- Table {t_idx} ({len(table.rows)} rows x {len(table.columns)} cols) ---\n")
        for r_idx, row in enumerate(table.rows):
            row_texts = [cell.text.strip().replace("\n", " ") for cell in row.cells]
            f.write(f"Row {r_idx}: {' | '.join(row_texts)}\n")

print(f"Extraction complete. Wrote to {output_txt}")
