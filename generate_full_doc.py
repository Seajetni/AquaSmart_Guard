# -*- coding: utf-8 -*-
"""
AquaSmart Guard Complete Academic Project Report Generator
Includes:
- Official Cover Page exactly matching shot.jpg (RMU Logo, Course, Instructor, Student)
- คำนำ (Preface)
- สารบัญ (Table of Contents) with dot leaders
- สารบัญตาราง (List of Tables)
- สารบัญภาพ (List of Figures)
- Chapters 1 to 5 (Full academic & engineering project report)
- บรรณานุกรม (References)
All formatted in Sarabun font with professional styling.
"""

import os
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

ASSETS_DIR = r"D:\aquarium-dashboard\doc_assets"
OUTPUT_DOCX = r"C:\Users\seaza\Downloads\AquaSmart_Guard.docx"

def set_run_font(run, font_name="Sarabun", size_pt=16, bold=False, italic=False, color_rgb=None):
    run.font.name = font_name
    run.font.size = Pt(size_pt)
    run.bold = bold
    run.italic = italic
    if color_rgb:
        run.font.color.rgb = color_rgb
    
    rPr = run._r.get_or_add_rPr()
    rFonts = OxmlElement('w:rFonts')
    rFonts.set(qn('w:ascii'), font_name)
    rFonts.set(qn('w:hAnsi'), font_name)
    rFonts.set(qn('w:cs'), font_name)
    rFonts.set(qn('w:eastAsia'), font_name)
    rPr.append(rFonts)

def add_p(doc, text="", font_size=16, bold=False, italic=False, align=WD_ALIGN_PARAGRAPH.LEFT,
          space_before=0, space_after=6, line_spacing=1.15, first_indent=None, color_rgb=None):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    if first_indent:
        p.paragraph_format.first_line_indent = Inches(first_indent)
    
    if text:
        run = p.add_run(text)
        set_run_font(run, "Sarabun", font_size, bold, italic, color_rgb)
    return p

def add_mixed_p(doc, runs_data, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=0, space_after=6,
                line_spacing=1.15, first_indent=None):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_before = Pt(space_before)
    p.paragraph_format.space_after = Pt(space_after)
    p.paragraph_format.line_spacing = line_spacing
    if first_indent:
        p.paragraph_format.first_line_indent = Inches(first_indent)
    
    for rdata in runs_data:
        text = rdata.get("text", "")
        size = rdata.get("size", 16)
        bold = rdata.get("bold", False)
        italic = rdata.get("italic", False)
        color = rdata.get("color", None)
        run = p.add_run(text)
        set_run_font(run, "Sarabun", size, bold, italic, color)
    return p

def add_chapter_title(doc, chapter_num, chapter_title):
    doc.add_page_break()
    p1 = add_p(doc, f"บทที่ {chapter_num}", font_size=20, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=14, space_after=6)
    p2 = add_p(doc, chapter_title, font_size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=18)
    return p1

def add_h1(doc, title):
    return add_p(doc, title, font_size=18, bold=True, space_before=14, space_after=6)

def add_h2(doc, title):
    return add_p(doc, title, font_size=16, bold=True, space_before=10, space_after=4)

def add_h3(doc, title):
    return add_p(doc, title, font_size=16, bold=True, space_before=8, space_after=3)

def add_bullet(doc, text, bold_prefix=None, indent_inches=0.4):
    p = doc.add_paragraph()
    p.paragraph_format.left_indent = Inches(indent_inches)
    p.paragraph_format.space_before = Pt(0)
    p.paragraph_format.space_after = Pt(4)
    p.paragraph_format.line_spacing = 1.15
    
    if bold_prefix:
        r1 = p.add_run(bold_prefix)
        set_run_font(r1, "Sarabun", 16, bold=True)
    r2 = p.add_run(text)
    set_run_font(r2, "Sarabun", 16, bold=False)
    return p

def add_toc_line(doc, title, page_str, level=1, bold=False):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(1)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.line_spacing = 1.15
    
    # Dot leader tab at 6.27 inches (the right margin of standard A4 with 1-in margins)
    p.paragraph_format.tab_stops.add_tab_stop(Inches(6.27), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    
    indent_space = "    " * (level - 1)
    r1 = p.add_run(f"{indent_space}{title}")
    set_run_font(r1, "Sarabun", 16, bold=bold)
    
    r2 = p.add_run(f"\t{page_str}")
    set_run_font(r2, "Sarabun", 16, bold=bold)
    return p

def add_image_box(doc, img_filename, caption, width_in=5.0):
    img_path = os.path.join(ASSETS_DIR, img_filename)
    if not os.path.exists(img_path):
        print(f"Warning: Image not found: {img_path}")
        return
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_before = Pt(8)
    p.paragraph_format.space_after = Pt(4)
    run = p.add_run()
    run.add_picture(img_path, width=Inches(width_in))
    
    cp = doc.add_paragraph()
    cp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    cp.paragraph_format.space_before = Pt(2)
    cp.paragraph_format.space_after = Pt(12)
    c_run = cp.add_run(caption)
    set_run_font(c_run, "Sarabun", 14, italic=True)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def set_cell_shading(cell, color_hex="F2F4F7"):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{color_hex}"/>')
    tcPr.append(shd)

def set_cell_border(cell, top="single", bottom="single", left="none", right="none", color="CCCCCC", sz="4"):
    tcPr = cell._tc.get_or_add_tcPr()
    borders = parse_xml(f'<w:tcBorders {nsdecls("w")}>'
                        f'<w:top w:val="{top}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
                        f'<w:left w:val="{left}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
                        f'<w:bottom w:val="{bottom}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
                        f'<w:right w:val="{right}" w:sz="{sz}" w:space="0" w:color="{color}"/>'
                        f'</w:tcBorders>')
    tcPr.append(borders)

def add_custom_table(doc, headers, data, col_widths=None, caption=None):
    if caption:
        cp = doc.add_paragraph()
        cp.paragraph_format.space_before = Pt(8)
        cp.paragraph_format.space_after = Pt(4)
        c_run = cp.add_run(caption)
        set_run_font(c_run, "Sarabun", 15, bold=True)
    
    rows = len(data) + 1
    cols = len(headers)
    table = doc.add_table(rows=rows, cols=cols)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    
    # Headers
    hdr_cells = table.rows[0].cells
    for i, title in enumerate(headers):
        hdr_cells[i].text = title
        set_cell_shading(hdr_cells[i], "0B2545")
        set_cell_margins(hdr_cells[i], top=120, bottom=120, left=150, right=150)
        set_cell_border(hdr_cells[i], top="single", bottom="single", left="none", right="none", color="0B2545", sz="8")
        p = hdr_cells[i].paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(0)
        if len(p.runs) > 0:
            set_run_font(p.runs[0], "Sarabun", 15, bold=True, color_rgb=RGBColor(255, 255, 255))
        hdr_cells[i].vertical_alignment = WD_ALIGN_VERTICAL.CENTER

    # Data Rows
    for r_idx, row_data in enumerate(data):
        row_cells = table.rows[r_idx + 1].cells
        bg_col = "F9FAFB" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            row_cells[c_idx].text = str(val)
            set_cell_shading(row_cells[c_idx], bg_col)
            set_cell_margins(row_cells[c_idx], top=80, bottom=80, left=150, right=150)
            set_cell_border(row_cells[c_idx], top="single", bottom="single", left="none", right="none", color="E5E7EB", sz="4")
            p = row_cells[c_idx].paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if (c_idx > 0 and len(str(val)) < 15) else WD_ALIGN_PARAGRAPH.LEFT
            p.paragraph_format.space_before = Pt(0)
            p.paragraph_format.space_after = Pt(0)
            if len(p.runs) > 0:
                set_run_font(p.runs[0], "Sarabun", 15, bold=False, color_rgb=RGBColor(30, 41, 59))
            row_cells[c_idx].vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            
    if col_widths:
        for r in table.rows:
            for c_idx, w in enumerate(col_widths):
                r.cells[c_idx].width = Inches(w)
                
    sp = doc.add_paragraph()
    sp.paragraph_format.space_before = Pt(0)
    sp.paragraph_format.space_after = Pt(8)
    return table

def build_aquasmart_document():
    doc = docx.Document()
    
    # Configure A4 & standard 1-inch margins
    for s in doc.sections:
        s.page_width = Inches(8.27)
        s.page_height = Inches(11.69)
        s.top_margin = Inches(1.0)
        s.bottom_margin = Inches(1.0)
        s.left_margin = Inches(1.0)
        s.right_margin = Inches(1.0)

    # =========================================================================
    # 1. หน้าปก (COVER PAGE - EXACTLY MATCHING shot.jpg)
    # =========================================================================
    # 1.1 University Logo (Rajabhat Maha Sarakham University)
    logo_path = os.path.join(ASSETS_DIR, "univ_logo_clean.png")
    if not os.path.exists(logo_path):
        logo_path = os.path.join(ASSETS_DIR, "univ_logo.png")
        
    p_logo = doc.add_paragraph()
    p_logo.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_logo.paragraph_format.space_before = Pt(12)
    p_logo.paragraph_format.space_after = Pt(36)
    if os.path.exists(logo_path):
        r_logo = p_logo.add_run()
        r_logo.add_picture(logo_path, width=Inches(1.35))
    
    # 1.2 Title block
    add_p(doc, "รายงาน", font_size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=6)
    add_p(doc, "เรื่อง ระบบตรวจวัดคุณภาพน้ำตู้ปลาอัจฉริยะ (AquaSmart Guard)", font_size=18, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=6)
    add_p(doc, "AquaSmart Guard: Smart Aquarium Water Quality Monitoring System with AI Assistant", font_size=15, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=42)
    
    # 1.3 Author block
    add_p(doc, "จัดทำโดย", font_size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=6)
    add_p(doc, "นายเจตนิพัทธ์ โชติเสวตร", font_size=16, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=4)
    add_p(doc, "รหัสนักศึกษา 6714631020", font_size=16, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=42)
    
    # 1.4 Instructor block
    add_p(doc, "อาจารย์ผู้สอน", font_size=16, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=6)
    add_p(doc, "อาจารย์ ดร.ปวริศ สารมะโน", font_size=16, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=42)
    
    # 1.5 Course & University block
    add_p(doc, "รายงานฉบับนี้เป็นส่วนหนึ่งของรายวิชาวิทยาการก้าวหน้าทางคอมพิวเตอร์", font_size=16, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=4)
    add_p(doc, "ภาคเรียนที่ 2 ปีการศึกษา 2567", font_size=16, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=4)
    add_p(doc, "สาขาวิชาคอมพิวเตอร์ศึกษา", font_size=16, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=4)
    add_p(doc, "คณะครุศาสตร์ มหาวิทยาลัยราชภัฏมหาสารคาม", font_size=16, bold=False, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=0, space_after=12)
    
    doc.add_page_break()

    # =========================================================================
    # 2. คำนำ (PREFACE)
    # =========================================================================
    add_p(doc, "คำนำ", font_size=20, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=14, space_after=18)
    
    add_p(doc, "รายงานฉบับนี้จัดทำขึ้นเพื่อนำเสนอผลการดำเนินงานโครงงานเรื่อง ระบบตรวจวัดคุณภาพน้ำตู้ปลาอัจฉริยะ (AquaSmart Guard: Smart Aquarium Water Quality Monitoring System with AI Assistant) ซึ่งเป็นส่วนหนึ่งของการศึกษารายวิชาวิทยาการก้าวหน้าทางคอมพิวเตอร์ โดยมีวัตถุประสงค์เพื่อประยุกต์ใช้เทคโนโลยีอินเทอร์เน็ตของสรรพสิ่ง (IoT) ร่วมกับไมโครคอนโทรลเลอร์ ESP32 เซนเซอร์ตรวจวัดคุณภาพน้ำ แพลตฟอร์มคลาวด์ Blynk IoT และเว็บแดชบอร์ด Next.js ตลอดจนการบูรณาการโมเดลปัญญาประดิษฐ์ Google Gemini AI ในการวิเคราะห์คุณภาพน้ำและให้คำแนะนำการดูแลปลาสวยงามตามสายพันธุ์อย่างชาญฉลาดและตรงจุด", first_indent=0.5)
    
    add_p(doc, "เนื้อหาภายในรายงานฉบับนี้แบ่งออกเป็น 5 บท ประกอบด้วย บทที่ 1 บทนำ แสดงถึงที่มาและความสำคัญ วัตถุประสงค์ ขอบเขต และประโยชน์ที่คาดว่าจะได้รับ, บทที่ 2 ทฤษฎีและงานวิจัยที่เกี่ยวข้อง อธิบายหลักการของพารามิเตอร์คุณภาพน้ำ เซนเซอร์ตรวจวัด ฮาร์ดแวร์ เทคโนโลยีเว็บ และปัญญาประดิษฐ์, บทที่ 3 วิธีการดำเนินงาน แสดงการออกแบบสถาปัตยกรรม การเชื่อมต่อวงจร และการพัฒนาซอฟต์แวร์ทั้งส่วนเฟิร์มแวร์และแดชบอร์ด, บทที่ 4 ผลการดำเนินงานและการทดลอง แสดงผลการทดสอบการเชื่อมต่อ การทดสอบระบบ AI การสอบเทียบความแม่นยำของเซนเซอร์ ตลอดจนการวิเคราะห์และแก้ไขปัญหาที่พบในการทดลองจริง และบทที่ 5 สรุปผลการดำเนินงาน ปัญหาอุปสรรค และข้อเสนอแนะในการพัฒนาต่อยอด", first_indent=0.5)
    
    add_p(doc, "ผู้จัดทำขอขอบพระคุณ อาจารย์ ดร.ปวริศ สารมะโน อาจารย์ผู้สอนประจำรายวิชา ที่ได้กรุณาถ่ายทอดองค์ความรู้ ให้คำแนะนำ แนวคิด คำปรึกษา ตลอดจนการตรวจทานและข้อเสนอแนะอันเป็นประโยชน์อย่างยิ่ง ทำให้โครงงานนี้สามารถดำเนินไปได้อย่างราบรื่นและประสบความสำเร็จลุล่วงด้วยดี", first_indent=0.5)
    
    add_p(doc, "ผู้จัดทำหวังเป็นอย่างยิ่งว่า รายงานฉบับนี้จะเป็นประโยชน์ต่อผู้เรียน นักศึกษา และผู้ที่สนใจศึกษาเกี่ยวกับการประยุกต์ใช้เทคโนโลยี IoT และปัญญาประดิษฐ์ หากมีข้อผิดพลาดหรือข้อบกพร่องประการใด ผู้จัดทำขอน้อมรับคำชี้แนะเพื่อนำไปปรับปรุงพัฒนาในโอกาสต่อไป", first_indent=0.5)
    
    # Signature
    add_p(doc, "ผู้จัดทำ", font_size=16, bold=False, align=WD_ALIGN_PARAGRAPH.RIGHT, space_before=24, space_after=4)
    add_p(doc, "นายเจตนิพัทธ์ โชติเสวตร", font_size=16, bold=False, align=WD_ALIGN_PARAGRAPH.RIGHT, space_before=0, space_after=4)
    
    doc.add_page_break()

    # =========================================================================
    # 3. สารบัญ (TABLE OF CONTENTS)
    # =========================================================================
    add_p(doc, "สารบัญ", font_size=20, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=14, space_after=14)
    
    # TOC Header row
    p_toc_hdr = doc.add_paragraph()
    p_toc_hdr.paragraph_format.tab_stops.add_tab_stop(Inches(6.27), WD_TAB_ALIGNMENT.RIGHT)
    r1 = p_toc_hdr.add_run("เรื่อง")
    set_run_font(r1, "Sarabun", 16, bold=True)
    r2 = p_toc_hdr.add_run("\tหน้า")
    set_run_font(r2, "Sarabun", 16, bold=True)
    
    # Preliminary
    add_toc_line(doc, "คำนำ", "ก", level=1, bold=True)
    add_toc_line(doc, "สารบัญ", "ข", level=1, bold=True)
    add_toc_line(doc, "สารบัญตาราง", "ค", level=1, bold=True)
    add_toc_line(doc, "สารบัญภาพ", "ง", level=1, bold=True)
    
    # Chapter 1
    add_toc_line(doc, "บทที่ 1 บทนำ", "1", level=1, bold=True)
    add_toc_line(doc, "1.1 ที่มาและความสำคัญของปัญหา", "1", level=2)
    add_toc_line(doc, "1.2 วัตถุประสงค์ของโครงงาน", "1", level=2)
    add_toc_line(doc, "1.3 ขอบเขตของโครงงาน", "2", level=2)
    add_toc_line(doc, "1.4 ประโยชน์ที่คาดว่าจะได้รับ", "3", level=2)
    add_toc_line(doc, "1.5 นิยามศัพท์เฉพาะ", "3", level=2)
    
    # Chapter 2
    add_toc_line(doc, "บทที่ 2 ทฤษฎีและงานวิจัยที่เกี่ยวข้อง", "4", level=1, bold=True)
    add_toc_line(doc, "2.1 คุณภาพน้ำและความสำคัญต่อการเลี้ยงปลาสวยงาม", "4", level=2)
    add_toc_line(doc, "2.2 อุปกรณ์และเซนเซอร์ตรวจวัดคุณภาพน้ำ", "5", level=2)
    add_toc_line(doc, "2.3 ไมโครคอนโทรลเลอร์และการสื่อสารไร้สาย ESP32", "6", level=2)
    add_toc_line(doc, "2.4 แพลตฟอร์มคลาวด์และ IoT (Blynk IoT Platform)", "6", level=2)
    add_toc_line(doc, "2.5 การพัฒนาเว็บแอปพลิเคชันยุคใหม่ (Web Technologies)", "7", level=2)
    add_toc_line(doc, "2.6 ปัญญาประดิษฐ์ Generative AI และ Google Gemini", "7", level=2)
    add_toc_line(doc, "2.7 งานวิจัยและโครงงานที่เกี่ยวข้อง", "8", level=2)
    
    # Chapter 3
    add_toc_line(doc, "บทที่ 3 วิธีการดำเนินงาน", "9", level=1, bold=True)
    add_toc_line(doc, "3.1 การวิเคราะห์และออกแบบระบบ", "9", level=2)
    add_toc_line(doc, "3.2 ขั้นตอนการดำเนินงาน", "9", level=2)
    add_toc_line(doc, "3.3 เครื่องมือที่ใช้", "10", level=2)
    add_toc_line(doc, "3.3.1 อุปกรณ์ Hardware", "10", level=3)
    add_toc_line(doc, "3.3.2 ซอฟต์แวร์ที่ใช้ (Software)", "14", level=3)
    add_toc_line(doc, "3.4 การเชื่อมต่อวงจร", "15", level=2)
    add_toc_line(doc, "3.5 โปรแกรม", "17", level=2)
    add_toc_line(doc, "3.5.1 โปรแกรมที่ใช้กับ ESP32", "17", level=3)
    add_toc_line(doc, "3.5.2 โปรแกรมที่ใช้กับ Dashboard", "18", level=3)
    
    # Chapter 4
    add_toc_line(doc, "บทที่ 4 ผลการดำเนินงานและการทดลอง", "19", level=1, bold=True)
    add_toc_line(doc, "4.1 ผลการประกอบวงจรและการติดตั้งระบบฮาร์ดแวร์ต้นแบบ", "19", level=2)
    add_toc_line(doc, "4.2 ผลการทำงานของระบบเฟิร์มแวร์และการส่งข้อมูล IoT ผ่าน Blynk Cloud", "19", level=2)
    add_toc_line(doc, "4.3 ผลการทำงานของ Web Dashboard", "20", level=2)
    add_toc_line(doc, "4.4 ผลการทดสอบระบบวิเคราะห์คุณภาพน้ำด้วย Google Gemini AI", "21", level=2)
    add_toc_line(doc, "4.5 ผลการทดสอบความแม่นยำและการสอบเทียบเซนเซอร์", "22", level=2)
    add_toc_line(doc, "4.6 การวิเคราะห์ปัญหาที่พบในการทดลองและแนวทางแก้ไข", "24", level=2)
    
    # Chapter 5
    add_toc_line(doc, "บทที่ 5 สรุปผลการดำเนินงานและข้อเสนอแนะ", "26", level=1, bold=True)
    add_toc_line(doc, "5.1 สรุปผลการดำเนินงาน", "26", level=2)
    add_toc_line(doc, "5.2 ปัญหาและข้อจำกัดของโครงงาน", "26", level=2)
    add_toc_line(doc, "5.3 ข้อเสนอแนะในการพัฒนาต่อยอด", "27", level=2)
    
    # References
    add_toc_line(doc, "บรรณานุกรม", "28", level=1, bold=True)
    
    doc.add_page_break()

    # =========================================================================
    # 4. สารบัญตาราง (LIST OF TABLES)
    # =========================================================================
    add_p(doc, "สารบัญตาราง", font_size=20, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=14, space_after=14)
    
    p_lot_hdr = doc.add_paragraph()
    p_lot_hdr.paragraph_format.tab_stops.add_tab_stop(Inches(6.27), WD_TAB_ALIGNMENT.RIGHT)
    r1 = p_lot_hdr.add_run("ตารางที่")
    set_run_font(r1, "Sarabun", 16, bold=True)
    r2 = p_lot_hdr.add_run("\tหน้า")
    set_run_font(r2, "Sarabun", 16, bold=True)
    
    add_toc_line(doc, "ตารางที่ 2.1 ค่ามาตรฐานพารามิเตอร์คุณภาพน้ำสำหรับสัตว์น้ำสวยงามยอดนิยม", "5")
    add_toc_line(doc, "ตารางที่ 3.1 การเชื่อมต่อโมดูลวัดค่า pH กับบอร์ด ESP32", "15")
    add_toc_line(doc, "ตารางที่ 3.2 การเชื่อมต่อโมดูลวัดค่า TDS กับบอร์ด ESP32", "15")
    add_toc_line(doc, "ตารางที่ 3.3 การเชื่อมต่อเซนเซอร์วัดอุณหภูมิ DS18B20 กับบอร์ด ESP32", "16")
    add_toc_line(doc, "ตารางที่ 4.1 ผลการทดสอบเสถียรภาพการรับส่งข้อมูลระหว่าง ESP32 และ Blynk Cloud", "20")
    add_toc_line(doc, "ตารางที่ 4.2 ตัวอย่างผลการทดสอบการวิเคราะห์คุณภาพน้ำและคำแนะนำจาก Google Gemini AI", "21")
    add_toc_line(doc, "ตารางที่ 4.3 ผลการสอบเทียบเซนเซอร์วัดค่า pH กับสารละลายบัฟเฟอร์มาตรฐาน 3 จุด", "22")
    add_toc_line(doc, "ตารางที่ 4.4 ผลการทดสอบเปรียบเทียบเซนเซอร์ TDS กับเครื่องวัดมาตรฐาน", "23")
    add_toc_line(doc, "ตารางที่ 4.5 ผลการทดสอบเปรียบเทียบเซนเซอร์วัดอุณหภูมิ DS18B20 กับเครื่องวัดมาตรฐาน", "24")
    
    doc.add_page_break()

    # =========================================================================
    # 5. สารบัญภาพ (LIST OF FIGURES)
    # =========================================================================
    add_p(doc, "สารบัญภาพ", font_size=20, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=14, space_after=14)
    
    p_lof_hdr = doc.add_paragraph()
    p_lof_hdr.paragraph_format.tab_stops.add_tab_stop(Inches(6.27), WD_TAB_ALIGNMENT.RIGHT)
    r1 = p_lof_hdr.add_run("รูปที่")
    set_run_font(r1, "Sarabun", 16, bold=True)
    r2 = p_lof_hdr.add_run("\tหน้า")
    set_run_font(r2, "Sarabun", 16, bold=True)
    
    add_toc_line(doc, "รูปที่ 3.1 บอร์ดไมโครคอนโทรลเลอร์ ESP32 DevKit", "10")
    add_toc_line(doc, "รูปที่ 3.2 หัววัดและโมดูลแปลงสัญญาณ pH Sensor (E-201-C)", "11")
    add_toc_line(doc, "รูปที่ 3.3 เซนเซอร์วัดค่าสารละลายในน้ำ TDS Sensor (Analog EC)", "12")
    add_toc_line(doc, "รูปที่ 3.4 เซนเซอร์วัดอุณหภูมิกันน้ำ DS18B20", "12")
    add_toc_line(doc, "รูปที่ 3.5 แผงต่อวงจรทดลอง Breadboard", "13")
    add_toc_line(doc, "รูปที่ 3.6 สายต่อวงจร Jumper Wire", "13")
    add_toc_line(doc, "รูปที่ 3.7 สายเชื่อมต่อ Micro-USB Cable", "14")
    add_toc_line(doc, "รูปที่ 3.8 ผังวงจรการเชื่อมต่อเซนเซอร์ตรวจวัดคุณภาพน้ำทั้งหมดกับบอร์ด ESP32", "16")

    # =========================================================================
    # 6. บทที่ 1 บทนำ
    # =========================================================================
    add_chapter_title(doc, "1", "บทนำ")
    
    add_h1(doc, "1.1 ที่มาและความสำคัญของปัญหา")
    add_p(doc, "การดูแลตู้ปลาให้ปลามีสุขภาพดีจำเป็นต้องทราบสภาพน้ำอย่างต่อเนื่อง แต่ผู้เลี้ยงส่วนใหญ่มักไม่ทราบว่าน้ำมีปัญหาจนกว่าปลาจะแสดงอาการผิดปกติ การตรวจน้ำด้วยชุดทดสอบเองต้องทำบ่อยและใช้เวลา อีกทั้งยังไม่มีเครื่องมือที่แสดงแนวโน้มค่าน้ำย้อนหลังให้เห็นภาพรวม และขาดคำแนะนำที่เฉพาะเจาะจงกับชนิดปลาที่เลี้ยงจริง ระบบที่พัฒนาขึ้นจึงนำเทคโนโลยี IoT มาใช้ในการตรวจวัดคุณภาพน้ำแบบเรียลไทม์ และส่งข้อมูลผ่านเครือข่าย Wi-Fi ขึ้น Cloud", first_indent=0.5)
    
    add_p(doc, "ระบบใช้ ESP32 เป็นอุปกรณ์หลักในการรับข้อมูลจากเซนเซอร์ 3 ตัว ได้แก่ pH Sensor (E-201-C) สำหรับวัดค่ากรด-ด่างของน้ำ TDS Sensor (Analog EC) สำหรับวัดค่าความสกปรกหรือสารละลายในน้ำ (PPM) และ DS18B20 สำหรับวัดอุณหภูมิน้ำแบบกันน้ำ จากนั้น ESP32 จะส่งข้อมูลขึ้น Blynk Cloud ผ่าน Virtual Pin ได้แก่ V0 (อุณหภูมิ) V1 (TDS) และ V2 (pH)", first_indent=0.5)
    
    add_p(doc, "ข้อมูลที่ส่งขึ้น Blynk Cloud จะถูกดึงมาแสดงผลผ่าน Dashboard ที่พัฒนาด้วย Next.js และ deploy บน Vercel ซึ่งช่วยให้ผู้ใช้งานดูค่าน้ำล่าสุดแบบเรียลไทม์ ดูกราฟแนวโน้มย้อนหลัง และเลือกหรือพิมพ์ชนิดปลาที่เลี้ยงเพื่อสอบถาม AI โดยระบบจะส่งค่าที่วัดได้พร้อมชนิดปลาไปให้ Google Gemini วิเคราะห์และแสดงคำแนะนำการดูแลบนหน้าเว็บ", first_indent=0.5)
    
    add_p(doc, "ดังนั้น ระบบ AquaSmart Guard จึงถูกพัฒนาขึ้นเพื่อให้ผู้เลี้ยงปลารู้ปัญหาก่อนที่ปลาจะป่วย เห็นแนวโน้มของคุณภาพน้ำย้อนหลัง เข้าถึงข้อมูลได้ทุกที่ผ่านเว็บ และได้รับคำแนะนำที่เฉพาะเจาะจงกับชนิดปลาที่เลี้ยงจริง", first_indent=0.5)
    
    add_h1(doc, "1.2 วัตถุประสงค์ของโครงงาน")
    add_bullet(doc, "เพื่อออกแบบและสร้างระบบตรวจวัดคุณภาพน้ำในตู้ปลาแบบเรียลไทม์ โดยใช้ไมโครคอนโทรลเลอร์ ESP32 ร่วมกับเซนเซอร์วัดค่าความเป็นกรด-ด่าง (pH), สารละลายรวมในน้ำ (TDS) และอุณหภูมิของน้ำ (Temperature)", bold_prefix="1. ")
    add_bullet(doc, "เพื่อเชื่อมต่อและส่งข้อมูลคุณภาพน้ำผ่านเครือข่าย Wi-Fi ขึ้นสู่ระบบคลาวด์ Blynk IoT ผ่านกลไก Virtual Pins (V0, V1, V2) แบบอัตโนมัติและต่อเนื่อง", bold_prefix="2. ")
    add_bullet(doc, "เพื่อพัฒนาเว็บแอปพลิเคชันแดชบอร์ด (Next.js Dashboard) สำหรับแสดงผลข้อมูลคุณภาพน้ำแบบเรียลไทม์ พร้อมกราฟแนวโน้มย้อนหลังหลายช่วงเวลา (นาที, ชั่วโมง, วัน, สัปดาห์, เดือน) และระบบจำลองสถานการณ์", bold_prefix="3. ")
    add_bullet(doc, "เพื่อประยุกต์ใช้โมเดลปัญญาประดิษฐ์ Google Gemini AI ในการวิเคราะห์คุณภาพน้ำ คำนวณคะแนนความสมบูรณ์ของน้ำ (Water Quality Score 0-100) และให้คำแนะนำการปรับสภาพน้ำที่เหมาะสมเฉพาะเจาะจงกับแต่ละสายพันธุ์ปลา", bold_prefix="4. ")
    add_bullet(doc, "เพื่อทดสอบประสิทธิภาพ ความแม่นยำของการสอบเทียบเซนเซอร์ และเสถียรภาพของระบบในการประยุกต์ใช้งานจริง", bold_prefix="5. ")
    
    add_h1(doc, "1.3 ขอบเขตของโครงงาน")
    add_p(doc, "โครงงานระบบ AquaSmart Guard ได้กำหนดขอบเขตการทำงานของระบบไว้ดังนี้")
    
    add_h2(doc, "1.3.1 ด้านอุปกรณ์ฮาร์ดแวร์และเซนเซอร์")
    add_bullet(doc, "ใช้ไมโครคอนโทรลเลอร์ ESP32 DevKit เป็นหน่วยประมวลผลหลักในการอ่านค่าและควบคุมระบบ", bold_prefix="1. ")
    add_bullet(doc, "ใช้เซนเซอร์ pH Sensor (E-201-C) สำหรับตรวจวัดค่าความเป็นกรด-ด่างของน้ำในช่วง 0.00 – 14.00 pH", bold_prefix="2. ")
    add_bullet(doc, "ใช้เซนเซอร์ TDS Sensor (Analog EC) สำหรับตรวจวัดสารละลายรวมและความสกปรกในน้ำในช่วง 0 – 1000 PPM", bold_prefix="3. ")
    add_bullet(doc, "ใช้เซนเซอร์ DS18B20 แบบกันน้ำ (Waterproof Probe) สำหรับตรวจวัดอุณหภูมิของน้ำในช่วง -55 ถึง 125 °C โดยมีความละเอียด 0.0625 °C", bold_prefix="4. ")
    
    add_h2(doc, "1.3.2 ด้านการสื่อสารและระบบคลาวด์")
    add_bullet(doc, "เชื่อมต่อเครือข่ายไร้สาย Wi-Fi ความถี่ 2.4 GHz เพื่อรับส่งข้อมูลระหว่าง ESP32 และเซิร์ฟเวอร์คลาวด์", bold_prefix="5. ")
    add_bullet(doc, "ส่งข้อมูลขึ้นแพลตฟอร์ม Blynk Cloud ผ่าน Virtual Pin ได้แก่ V0 (อุณหภูมิ), V1 (TDS) และ V2 (pH)", bold_prefix="6. ")
    add_bullet(doc, "กำหนดรอบเวลาการอ่านและส่งข้อมูลจากเซนเซอร์ทุก 3 วินาที เพื่อให้ข้อมูลมีความสดใหม่และเป็นเรียลไทม์", bold_prefix="7. ")
    
    add_h2(doc, "1.3.3 ด้านเว็บแอปพลิเคชันและการแสดงผล")
    add_bullet(doc, "พัฒนา Web Dashboard ด้วยเฟรมเวิร์ก Next.js 14 และเผยแพร่ระบบบน Vercel (https://aqua-smart-guard.vercel.app)", bold_prefix="8. ")
    add_bullet(doc, "แสดงค่าน้ำล่าสุด ได้แก่ pH, TDS และอุณหภูมิ แบบเรียลไทม์ พร้อมตัวแสดงสถานะไฟกระพริบ (Live Pulse Indicator)", bold_prefix="9. ")
    add_bullet(doc, "แสดงผลกราฟแนวโน้มค่าวัดย้อนหลังแบบ 3 แกน (Multi-axis Chart) ด้วย Chart.js พร้อมตัวเลือกช่วงเวลา (นาที, ชั่วโมง, วัน, สัปดาห์, เดือน)", bold_prefix="10. ")
    add_bullet(doc, "มีโหมดจำลองสถานการณ์ (Simulation & Manual Testing) เพื่อทดสอบการแจ้งเตือนและการทำงานของระบบโดยไม่ต้องใช้เซนเซอร์จริง", bold_prefix="11. ")
    
    add_h2(doc, "1.3.4 ด้านปัญญาประดิษฐ์และการประมวลผล")
    add_bullet(doc, "ผู้ใช้สามารถเลือกสายพันธุ์ปลายอดนิยม (เช่น ปลาหมอสี, ปลาเทวดา, ปลาคาร์ฟ, ปลาปอมปาดัวร์, ปลากัด, ปลาหางนกยูง, ปลามังกร, กุ้งแคระ ฯลฯ) หรือพิมพ์ค้นหาสายพันธุ์เพิ่มเติมได้", bold_prefix="12. ")
    add_bullet(doc, "ใช้ Google Gemini AI (โมเดล gemini-3.5-flash-lite / Flash) ในการวิเคราะห์คุณภาพน้ำเทียบกับความต้องการของสายพันธุ์ปลา และแสดงคำแนะนำการดูแลบนหน้าเว็บ", bold_prefix="13. ")
    add_bullet(doc, "มีระบบคำนวณคะแนนความสมบูรณ์ของน้ำ (Water Quality Score 0-100) และระบบฐานความรู้สำรอง (Fallback Knowledge Base) ในกรณีเครือข่ายขัดข้อง", bold_prefix="14. ")
    
    add_h2(doc, "1.3.5 ขอบเขตและข้อจำกัดของระบบ")
    add_bullet(doc, "ระบบเป็นระบบสำหรับติดตาม ตรวจวัด และให้คำแนะนำเบื้องต้นเพื่อประกอบการตัดสินใจของผู้เลี้ยงปลา มิใช่ระบบบำบัดน้ำหรือจ่ายยาอัตโนมัติเต็มรูปแบบ", bold_prefix="15. ")
    
    add_h1(doc, "1.4 ประโยชน์ที่คาดว่าจะได้รับ")
    add_bullet(doc, "รู้ปัญหาก่อนปลาป่วย: สามารถตรวจจับความเปลี่ยนแปลงของคุณภาพน้ำที่ผิดปกติได้ทันทีโดยไม่ต้องรอให้ปลาแสดงอาการป่วยหรือตาย ช่วยลดความสูญเสีย", bold_prefix="1. ")
    add_bullet(doc, "เห็นแนวโน้มย้อนหลัง: กราฟแสดงค่าน้ำย้อนหลังช่วยให้ผู้เลี้ยงเข้าใจพฤติกรรมการเปลี่ยนแปลงของระบบกรองและวงจรชีวภาพในตู้ปลาได้อย่างชัดเจน", bold_prefix="2. ")
    add_bullet(doc, "เข้าถึงได้ทุกที่ทุกเวลา: ผู้เลี้ยงสามารถติดตามสภาวะตู้ปลาผ่าน Web Dashboard ได้อย่างสะดวกรวดเร็วทั้งบนสมาร์ตโฟน แท็บเล็ต และคอมพิวเตอร์", bold_prefix="3. ")
    add_bullet(doc, "คำแนะนำเฉพาะทางตามสายพันธุ์: ได้รับคำแนะนำและวิธีปรับสภาพน้ำที่ถูกต้องเหมาะสมกับสายพันธุ์ปลาที่เลี้ยงจริงผ่านการประมวลผลของปัญญาประดิษฐ์", bold_prefix="4. ")
    add_bullet(doc, "ต้นแบบเทคโนโลยีอัจฉริยะราคาประหยัด: ได้ระบบต้นแบบ IoT และแดชบอร์ดที่สามารถนำไปต่อยอดใช้ในฟาร์มเพาะพันธุ์สัตว์น้ำหรือธุรกิจปลาสวยงามได้จริง", bold_prefix="5. ")
    
    add_h1(doc, "1.5 นิยามศัพท์เฉพาะ")
    add_bullet(doc, "Internet of Things (IoT): เทคโนโลยีการเชื่อมต่ออุปกรณ์ฮาร์ดแวร์เข้ากับเครือข่ายอินเทอร์เน็ต เพื่อการแลกเปลี่ยนและควบคุมข้อมูลทางไกล", bold_prefix="• ")
    add_bullet(doc, "ESP32: ไมโครคอนโทรลเลอร์ขนาด 32-bit ที่มีโมดูล Wi-Fi และ Bluetooth ในตัว เหมาะสำหรับงานประมวลผล IoT", bold_prefix="• ")
    add_bullet(doc, "pH (Potential of Hydrogen): ดัชนีชี้วัดความเป็นกรดหรือด่างของสารละลาย มีค่าตั้งแต่ 0 (กรดเข้มข้น) ถึง 14 (ด่างเข้มข้น) โดย 7 คือค่าเป็นกลาง", bold_prefix="• ")
    add_bullet(doc, "TDS (Total Dissolved Solids): ปริมาณของแข็งหรือสารละลายทั้งหมดที่ละลายอยู่ในน้ำ มีหน่วยวัดเป็น PPM (Parts Per Million)", bold_prefix="• ")
    add_bullet(doc, "Virtual Pin: ขาพินเสมือนที่สร้างขึ้นบนแพลตฟอร์ม Blynk เพื่อใช้ในการแลกเปลี่ยนข้อมูลระหว่างฮาร์ดแวร์กับระบบคลาวด์โดยไม่ยึดติดกับขาจริง", bold_prefix="• ")
    add_bullet(doc, "Water Quality Score: คะแนนดัชนีชี้วัดความเหมาะสมของน้ำตั้งแต่ 0 ถึง 100 คะแนน คำนวณจากระยะห่างของค่าวัดจริงเทียบกับเกณฑ์มาตรฐานของปลาแต่ละสายพันธุ์", bold_prefix="• ")

    # =========================================================================
    # 7. บทที่ 2 ทฤษฎีและงานวิจัยที่เกี่ยวข้อง
    # =========================================================================
    add_chapter_title(doc, "2", "ทฤษฎีและงานวิจัยที่เกี่ยวข้อง")
    
    add_h1(doc, "2.1 คุณภาพน้ำและความสำคัญต่อการเลี้ยงปลาสวยงาม")
    add_p(doc, "คุณภาพน้ำถือเป็นปัจจัยสำคัญที่สุดในการดำรงชีวิตของปลาและสัตว์น้ำ เนื่องจากน้ำเป็นทั้งสิ่งแวดล้อม แหล่งออกซิเจน และแหล่งขับถ่ายของเสีย การเปลี่ยนแปลงของค่าน้ำแม้เพียงเล็กน้อยอาจส่งผลต่อความเครียด การเจริญเติบโต และภูมิคุ้มกันของปลา พารามิเตอร์หลักที่ระบบ AquaSmart Guard ให้ความสำคัญประกอบด้วย", first_indent=0.5)
    
    add_h2(doc, "2.1.1 ค่าความเป็นกรด-ด่าง (pH)")
    add_p(doc, "ค่า pH เป็นตัวชี้วัดความเข้มข้นของไฮโดรเจนไอออนในน้ำ ปลาสวยงามน้ำจืดส่วนใหญ่จะอาศัยอยู่ในช่วง pH ระหว่าง 6.5 ถึง 8.5 หากค่า pH ต่ำกว่า 6.0 (สภาพกรด) จะทำลายเยื่อบุผิวเหงือก ขัดขวางการแลกเปลี่ยนก๊าซ และทำให้ปลาผลิตเมือกมากผิดปกติ ในทางตรงกันข้าม หากค่า pH สูงเกิน 8.5 (สภาพด่าง) สารประกอบแอมโมเนียมที่ไม่เป็นพิษ (NH4+) จะเปลี่ยนรูปกลายเป็นแอมโมเนียอิสระ (NH3) ซึ่งมีพิษร้ายแรงต่อระบบประสาทและการหายใจของปลาอย่างเฉียบพลัน", first_indent=0.5)
    
    add_h2(doc, "2.1.2 ค่าสารละลายรวมในน้ำ (TDS / Total Dissolved Solids)")
    add_p(doc, "ค่า TDS สะท้อนถึงปริมาณแร่ธาตุ เกลือ และสารอินทรีย์-อนินทรีย์ที่ละลายอยู่ในน้ำ มีหน่วยเป็นมิลลิกรัมต่อลิตร (mg/L) หรือส่วนในล้านส่วน (PPM) ค่า TDS มีผลโดยตรงต่อความดันออสโมซิส (Osmotic Pressure) ในเซลล์ของปลา หากน้ำมี TDS ต่ำหรือสูงเกินกว่าธรรมชาติของสายพันธุ์นั้นๆ ปลาจะต้องใช้พลังงานสูงมากในการควบคุมสมดุลน้ำและเกลือแร่ (Osmoregulation) นอกจากนี้ การเพิ่มขึ้นอย่างรวดเร็วของ TDS มักเป็นสัญญาณของการสะสมของเสีย การให้อาหารเกิน และการขาดการเปลี่ยนถ่ายน้ำ", first_indent=0.5)
    
    add_h2(doc, "2.1.3 อุณหภูมิของน้ำ (Temperature)")
    add_p(doc, "ปลาเป็นสัตว์เลือดเย็น (Poikilothermic animals) อุณหภูมิร่างกายจึงเปลี่ยนแปลงตามอุณหภูมิของน้ำโดยตรง อุณหภูมิที่เหมาะสมสำหรับปลาสวยงามเขตร้อนส่วนใหญ่จะอยู่ในช่วง 25.0 – 28.5 °C อุณหภูมิที่สูงขึ้นจะทำให้อัตราเมแทบอลิซึมของปลาเร็วขึ้น ต้องการออกซิเจนมากขึ้น แต่ในขณะเดียวกันความสามารถในการละลายของก๊าซออกซิเจนในน้ำ (DO) จะลดลง ส่วนอุณหภูมิที่ลดต่ำลงหรือแกว่งตัวเกิน 2 °C ในรอบวัน จะกดภูมิคุ้มกันของปลา ทำให้ปลาเป็นโรคจุดขาว (Ich) และติดเชื้อแบคทีเรียได้ง่าย", first_indent=0.5)
    
    # Table 2.1
    fish_headers = ["สายพันธุ์ปลา", "pH ที่เหมาะสม", "TDS (PPM)", "อุณหภูมิ (°C)", "ลักษณะน้ำที่ชอบ"]
    fish_data = [
        ["ปลาหมอสี (Cichlid)", "7.5 – 8.6", "250 – 450", "25.0 – 28.5", "น้ำด่างปานกลาง แร่ธาตุสูง"],
        ["ปลาปอมปาดัวร์ (Discus)", "6.0 – 6.8", "50 – 150", "28.0 – 30.5", "น้ำอ่อน สะอาดมาก นุ่ม"],
        ["ปลาเทวดา (Angelfish)", "6.5 – 7.2", "100 – 200", "26.0 – 29.0", "น้ำอ่อนถึงปานกลาง กรดอ่อน"],
        ["ปลากัด (Betta)", "6.5 – 7.5", "100 – 300", "24.0 – 28.0", "น้ำนิ่ง ค่าเป็นกลาง"],
        ["ปลาหางนกยูง (Guppy)", "7.0 – 8.0", "150 – 400", "22.0 – 28.0", "น้ำด่างอ่อน แร่ธาตุปานกลาง"],
        ["ปลาคาร์ฟ / ปลาทอง", "7.2 – 7.8", "150 – 350", "20.0 – 26.0", "น้ำสะอาด ออกซิเจนสูง"],
        ["ปลามังกร (Arowana)", "6.5 – 7.5", "100 – 250", "26.0 – 30.0", "น้ำสะอาด การไหลเวียนดี"],
        ["กุ้งแคระ (Dwarf Shrimp)", "6.8 – 7.6", "120 – 200", "22.0 – 26.0", "น้ำเสถียร ไม่ชอบการแกว่ง"]
    ]
    add_custom_table(doc, fish_headers, fish_data, col_widths=[1.5, 1.1, 1.1, 1.1, 1.7], caption="ตารางที่ 2.1 ค่ามาตรฐานพารามิเตอร์คุณภาพน้ำสำหรับสัตว์น้ำสวยงามยอดนิยม")

    add_h1(doc, "2.2 อุปกรณ์และเซนเซอร์ตรวจวัดคุณภาพน้ำ")
    add_p(doc, "การตรวจวัดคุณภาพน้ำอย่างแม่นยำและต่อเนื่องจำเป็นต้องอาศัยหลักการทางเคมีไฟฟ้าและเซนเซอร์ที่เชื่อถือได้", first_indent=0.5)
    
    add_h2(doc, "2.2.1 เซนเซอร์วัดค่าความเป็นกรด-ด่าง (pH Sensor E-201-C)")
    add_p(doc, "โพรบวัด pH รุ่น E-201-C เป็นอิเล็กโทรดแก้วรวม (Combination Glass Electrode) ประกอบด้วย Measuring Electrode ที่มีเยื่อแก้วบางไวต่อไฮโดรเจนไอออน และ Reference Electrode (Ag/AgCl) ที่บรรจุสารละลายอิเล็กโทรไลต์ KCl ศักย์ไฟฟ้าที่เกิดขึ้นระหว่างผิวแก้วกับของเหลวจะแปรผันตรงกับค่า pH ตามสมการ Nernst โมดูลแปลงสัญญาณจะแปลงแรงดันไฟฟ้าระดับมิลลิโวลต์ (mV) ให้เป็นสัญญาณแรงดันแอนะล็อกช่วง 0 – 3.3V / 5V ที่ ESP32 สามารถอ่านผ่าน ADC ได้อย่างแม่นยำ", first_indent=0.5)
    
    add_h2(doc, "2.2.2 เซนเซอร์วัดค่าสารละลายรวม (Analog TDS Sensor)")
    add_p(doc, "เซนเซอร์ TDS ตรวจวัดการนำไฟฟ้า (Electrical Conductivity: EC) ของของเหลวโดยใช้ขั้วไฟฟ้าโลหะคู่ที่ป้อนสัญญาณกระแสสลับ (AC Signal) เพื่อป้องกันปรากฏการณ์อิเล็กโทรไลซิสและการเกิดฟองก๊าซสะสมที่ขั้วโพรบ (Electrochemical Polarization) จากนั้นจะแปลงค่าการนำไฟฟ้าเป็นแรงดันแอนะล็อก และใช้สูตรแปลงค่า EC เป็น PPM โดยคำนึงถึงค่าชดเชยอุณหภูมิ (Temperature Compensation Coefficient)", first_indent=0.5)
    
    add_h2(doc, "2.2.3 เซนเซอร์วัดอุณหภูมิดิจิทัลกันน้ำ (DS18B20 Waterproof)")
    add_p(doc, "DS18B20 เป็นเซนเซอร์อุณหภูมิดิจิทัลที่มีความแม่นยำสูง บรรจุอยู่ในปลอกสเตนเลสกันน้ำ สื่อสารผ่านโปรโตคอล 1-Wire Bus ซึ่งใช้สายสัญญาณเพียงเส้นเดียวในการแลกเปลี่ยนข้อมูลกับ ESP32 โดยอาศัยตัวต้านทาน Pull-up ขนาด 4.7 kΩ เซนเซอร์สามารถวัดอุณหภูมิได้ตั้งแต่ -55 ถึง +125 °C โดยมีค่าความคลาดเคลื่อนเพียง ±0.5 °C ในช่วง -10 ถึง +85 °C และส่งข้อมูลเป็นตัวเลขดิจิทัลโดยตรง จึงไม่ถูกลดทอนสัญญาณจากความยาวของสาย", first_indent=0.5)

    add_h1(doc, "2.3 ไมโครคอนโทรลเลอร์และการสื่อสารไร้สาย ESP32")
    add_p(doc, "ESP32 DevKit เป็นบอร์ดไมโครคอนโทรลเลอร์ประสิทธิภาพสูง ผลิตโดย Espressif Systems ใช้สถาปัตยกรรม Xtensa Dual-Core 32-bit LX6 ความเร็วสัญญาณนาฬิกาสูงสุด 240 MHz มีหน่วยความจำ SRAM ขนาด 520 KB และ Flash Memory ขนาด 4 MB บอร์ดมีวงจร Wi-Fi 802.11 b/g/n และ Bluetooth 4.2 BR/EDR/BLE ในตัว มีพอร์ต Analog-to-Digital Converter (ADC) ความละเอียด 12-bit (0 – 4095) ซึ่งช่วยให้อ่านสัญญาณแอนะล็อกของเซนเซอร์ pH และ TDS ได้อย่างละเอียด", first_indent=0.5)

    add_h1(doc, "2.4 แพลตฟอร์มคลาวด์และ IoT (Blynk IoT Platform)")
    add_p(doc, "Blynk IoT เป็นแพลตฟอร์มสำหรับเชื่อมต่ออุปกรณ์ฮาร์ดแวร์เข้าสู่ระบบคลาวด์ที่ได้รับความนิยมสูง มีระบบรักษาความปลอดภัยด้วย Auth Token ประจำอุปกรณ์ และใช้แนวคิด Virtual Pins ในการแมปข้อมูลระหว่างโค้ดโปรแกรมกับส่วนแสดงผล การใช้ Virtual Pins ช่วยให้อุปกรณ์ ESP32 สามารถอัปเดตข้อมูลขึ้นคลาวด์ได้โดยตรง และแอปพลิเคชันภายนอกสามารถดึงข้อมูลผ่าน Blynk REST API ในรูปแบบ JSON ได้ทันที", first_indent=0.5)

    add_h1(doc, "2.5 การพัฒนาเว็บแอปพลิเคชันยุคใหม่ (Web Technologies)")
    add_p(doc, "เพื่อให้แดชบอร์ดทำงานได้รวดเร็ว ตอบสนองดี และมีความสวยงาม โครงงานจึงเลือกใช้เทคโนโลยีเว็บระดับแนวหน้า ได้แก่", first_indent=0.5)
    add_bullet(doc, "Next.js 14 App Router: สถาปัตยกรรมเว็บสมัยใหม่ที่รองรับ Server Components, Dynamic Route Caching และ Serverless API Routes บนโฮสติ้ง Vercel", bold_prefix="• ")
    add_bullet(doc, "React 18 & Tailwind CSS: สำหรับสร้าง Modern Glassmorphism UI พร้อมเอฟเฟกต์แอนิเมชันและตัวบ่งชี้สถานะ Live Pulse แบบเรียลไทม์", bold_prefix="• ")
    add_bullet(doc, "Chart.js Multi-axis: กราฟเส้นแบบ 3 แกนอิสระที่สามารถเปรียบเทียบ pH, TDS และ อุณหภูมิบนแกนเวลาเดียวกันได้อย่างชัดเจน พร้อมตัวกรองช่วงเวลาย้อนหลัง (นาที, ชั่วโมง, วัน, สัปดาห์, เดือน)", bold_prefix="• ")

    add_h1(doc, "2.6 ปัญญาประดิษฐ์ Generative AI และ Google Gemini")
    add_p(doc, "Google Gemini เป็นโมเดลภาษาขนาดใหญ่ (LLM) ที่มีความสามารถในการให้เหตุผลเชิงลึกและการวิเคราะห์ข้อมูลหลายมิติ ในโครงงานนี้เลือกใช้โมเดล gemini-3.5-flash-lite / Flash ซึ่งให้ความเร็วในการตอบสนองสูง (Low Latency) และมีความแม่นยำในการทำความเข้าใจสัตวศาสตร์ทางน้ำ การป้อนข้อมูลบริบท (Prompt Engineering) ที่ประกอบด้วย ชนิดปลา, ค่าวัดจริง, ค่ามาตรฐานอ้างอิง และประวัติการเปลี่ยนแปลง ช่วยให้ Gemini สร้างคำแนะนำการปรับสภาพน้ำที่ถูกต้อง ปลอดภัย และตรงจุดที่สุด พร้อมกันนี้ระบบยังมี Fallback Knowledge Base เก็บฐานข้อมูลออฟไลน์ไว้ป้องกันข้อผิดพลาดกรณีเครือข่ายขัดข้อง", first_indent=0.5)

    add_h1(doc, "2.7 งานวิจัยและโครงงานที่เกี่ยวข้อง")
    add_p(doc, "จากการทบทวนวรรณกรรมที่เกี่ยวข้องกับการตรวจวัดคุณภาพน้ำในอดีต พบว่าระบบส่วนใหญ่จะเน้นเพียงการอ่านค่าและแสดงผลบนจอ LCD หรือแจ้งเตือนพื้นฐานเมื่อค่าเกินขีดจำกัด แต่ยังขาดความยืดหยุ่นในการปรับแต่งตามชนิดของปลาที่เลี้ยง และไม่มีระบบปัญญาประดิษฐ์ที่ช่วยวิเคราะห์เชิงลึก ระบบ AquaSmart Guard จึงถูกออกแบบมาเพื่อเติมเต็มช่องว่างดังกล่าว โดยผสาน IoT, กราฟแนวโน้มพหุแกน และผู้ช่วย AI ไว้ในระบบเดียวอย่างสมบูรณ์", first_indent=0.5)

    # =========================================================================
    # 8. บทที่ 3 วิธีการดำเนินงาน
    # =========================================================================
    add_chapter_title(doc, "3", "วิธีการดำเนินงาน")
    
    add_h1(doc, "3.1 การวิเคราะห์และออกแบบระบบ")
    add_p(doc, "สถาปัตยกรรมของระบบ AquaSmart Guard แบ่งออกเป็น 4 ส่วนหลัก ได้แก่ ส่วนตรวจวัดและประมวลผลฮาร์ดแวร์ (ESP32), ส่วนระบบคลาวด์และจัดการข้อมูล (Blynk IoT), ส่วนแดชบอร์ดแสดงผล (Next.js บน Vercel) และส่วนบริการวิเคราะห์อัจฉริยะ (Google Gemini AI)", first_indent=0.5)
    
    add_p(doc, "โครงสร้างการไหลของข้อมูลเริ่มต้นจากเซนเซอร์ตรวจวัดค่าน้ำ -> ส่งค่าแอนะล็อกและดิจิทัลเข้าสู่บอร์ด ESP32 -> บอร์ดทำการแปลงหน่วยและกรองสัญญาณรบกวน -> ส่งข้อมูลผ่าน Wi-Fi ไปยัง Blynk Cloud ที่ Virtual Pin V0 (อุณหภูมิ), V1 (TDS) และ V2 (pH) -> Web Dashboard ดึงข้อมูลผ่าน REST API ทุก 3 วินาที -> แสดงผลแบบเรียลไทม์ บันทึกกราฟย้อนหลัง -> เมื่อผู้ใช้ระบุชนิดปลาและกดวิเคราะห์ ระบบจะรวบรวมค่าน้ำส่งให้ Google Gemini เพื่อคำนวณ Water Quality Score และแสดงข้อแนะนำการดูแลตู้ปลา", first_indent=0.5)
    
    add_h1(doc, "3.2 ขั้นตอนการดำเนินงาน")
    add_bullet(doc, "ขั้นที่ 1 ศึกษาและรวบรวมข้อมูล: ศึกษาพารามิเตอร์คุณภาพน้ำที่เหมาะสมสำหรับปลาสวยงามแต่ละชนิด และศึกษาคู่มือการใช้งานของเซนเซอร์", bold_prefix="1. ")
    add_bullet(doc, "ขั้นที่ 2 ออกแบบและจัดหาอุปกรณ์ฮาร์ดแวร์: จัดเตรียมบอร์ด ESP32, pH Sensor, TDS Sensor, DS18B20 และอุปกรณ์ประกอบวงจร", bold_prefix="2. ")
    add_bullet(doc, "ขั้นที่ 3 พัฒนาเฟิร์มแวร์ ESP32: พัฒนาโปรแกรมอ่านค่าเซนเซอร์ สอบเทียบสัญญาณ และส่งข้อมูลขึ้น Blynk Cloud แบบ Non-blocking", bold_prefix="3. ")
    add_bullet(doc, "ขั้นที่ 4 พัฒนาเว็บแดชบอร์ด (Web Dashboard): สร้างส่วนติดต่อผู้ใช้ด้วย Next.js, ออกแบบกราฟ Multi-axis ด้วย Chart.js และสร้างระบบจำลองค่า", bold_prefix="4. ")
    add_bullet(doc, "ขั้นที่ 5 เชื่อมต่อและพัฒนาผู้ช่วย AI: ออกแบบ API Route และ Prompt Template ในการเรียก Google Gemini เพื่อวิเคราะห์คุณภาพน้ำและให้คะแนน", bold_prefix="5. ")
    add_bullet(doc, "ขั้นที่ 6 ทดสอบ สอบเทียบ และปรับปรุงระบบ: ทดสอบความแม่นยำของเซนเซอร์กับสารละลายมาตรฐาน และทดสอบการทำงานต่อเนื่องของระบบ", bold_prefix="6. ")

    add_h1(doc, "3.3 เครื่องมือที่ใช้")
    add_p(doc, "เครื่องมือที่ใช้ในการพัฒนาระบบแบ่งออกเป็น อุปกรณ์ฮาร์ดแวร์ และซอฟต์แวร์ ดังนี้")
    
    add_h2(doc, "3.3.1 อุปกรณ์ Hardware")
    
    add_p(doc, "1. ESP32 DevKit", font_size=16, bold=True)
    add_image_box(doc, "image8.jpg", "รูปที่ 3.1 บอร์ดไมโครคอนโทรลเลอร์ ESP32 DevKit", width_in=3.2)
    add_mixed_p(doc, [
        {"text": "หน้าที่: ", "size": 16, "bold": True},
        {"text": "เป็นไมโครคอนโทรลเลอร์หลักของระบบ ทำหน้าที่อ่านค่าจากเซนเซอร์ pH, TDS และอุณหภูมิ เชื่อมต่อ Wi-Fi และส่งข้อมูลขึ้น Blynk Cloud ผ่าน Virtual Pin", "size": 16, "bold": False}
    ], space_after=10)
    
    add_p(doc, "2. pH Sensor (E-201-C)", font_size=16, bold=True)
    add_image_box(doc, "image6.png", "รูปที่ 3.2 หัววัดและโมดูลแปลงสัญญาณ pH Sensor (E-201-C)", width_in=4.8)
    add_mixed_p(doc, [
        {"text": "หน้าที่: ", "size": 16, "bold": True},
        {"text": "ใช้สำหรับตรวจวัดค่ากรด-ด่าง (pH) ของน้ำในตู้ปลา โดยแปลงศักย์ไฟฟ้าเคมีเป็นสัญญาณแอนะล็อกส่งไปยัง ESP32 ที่ขา GPIO 35 และส่งขึ้น Blynk ที่ Virtual Pin V2", "size": 16, "bold": False}
    ], space_after=10)
    
    add_p(doc, "3. TDS Sensor (Analog EC)", font_size=16, bold=True)
    add_image_box(doc, "image5.png", "รูปที่ 3.3 เซนเซอร์วัดค่าสารละลายในน้ำ TDS Sensor (Analog EC)", width_in=4.8)
    add_mixed_p(doc, [
        {"text": "หน้าที่: ", "size": 16, "bold": True},
        {"text": "ใช้สำหรับตรวจวัดค่าความสกปรกหรือสารละลายรวมในน้ำ (TDS) หน่วย PPM โดยส่งสัญญาณแอนะล็อกไปยัง ESP32 ที่ขา GPIO 34 และส่งขึ้น Blynk ที่ Virtual Pin V1", "size": 16, "bold": False}
    ], space_after=10)
    
    add_p(doc, "4. DS18B20 Waterproof Probe", font_size=16, bold=True)
    add_image_box(doc, "image4.png", "รูปที่ 3.4 เซนเซอร์วัดอุณหภูมิกันน้ำ DS18B20", width_in=4.8)
    add_mixed_p(doc, [
        {"text": "หน้าที่: ", "size": 16, "bold": True},
        {"text": "ใช้สำหรับตรวจวัดอุณหภูมิน้ำแบบกันน้ำ (Waterproof) สื่อสารผ่านโปรโตคอล 1-Wire เข้าสู่ขา GPIO 32 ของ ESP32 และส่งค่าขึ้น Blynk ที่ Virtual Pin V0", "size": 16, "bold": False}
    ], space_after=10)
    
    add_p(doc, "5. Breadboard", font_size=16, bold=True)
    add_image_box(doc, "image1.jpg", "รูปที่ 3.5 แผงต่อวงจรทดลอง Breadboard", width_in=3.4)
    add_mixed_p(doc, [
        {"text": "หน้าที่: ", "size": 16, "bold": True},
        {"text": "ใช้สำหรับทดลองและประกอบวงจร โดยสามารถเสียบอุปกรณ์และสาย Jumper เพื่อทดสอบการทำงานของระบบโดยไม่ต้องบัดกรีวงจร", "size": 16, "bold": False}
    ], space_after=10)
    
    add_p(doc, "6. Jumper Wire", font_size=16, bold=True)
    add_image_box(doc, "image2.jpg", "รูปที่ 3.6 สายต่อวงจร Jumper Wire", width_in=3.0)
    add_mixed_p(doc, [
        {"text": "หน้าที่: ", "size": 16, "bold": True},
        {"text": "ใช้สำหรับเชื่อมต่อขา ESP32 กับเซนเซอร์ต่าง ๆ บน Breadboard โดยสาย Jumper เป็นอุปกรณ์พื้นฐานที่ใช้ในการสร้างวงจรต้นแบบ", "size": 16, "bold": False}
    ], space_after=10)
    
    add_p(doc, "7. USB Cable", font_size=16, bold=True)
    add_image_box(doc, "image3.png", "รูปที่ 3.7 สายเชื่อมต่อ Micro-USB Cable", width_in=4.8)
    add_mixed_p(doc, [
        {"text": "หน้าที่: ", "size": 16, "bold": True},
        {"text": "ใช้สำหรับเชื่อมต่อ ESP32 กับคอมพิวเตอร์ เพื่อจ่ายไฟ อัปโหลดโปรแกรม และตรวจสอบข้อมูลผ่าน Serial Monitor โดยชุดพัฒนา ESP32 สามารถใช้ USB สำหรับการเขียนโปรแกรมและจ่ายไฟได้", "size": 16, "bold": False}
    ], space_after=10)

    add_h2(doc, "3.3.2 ซอฟต์แวร์ที่ใช้ (Software)")
    add_bullet(doc, "Arduino IDE 2.x: โปรแกรมสำหรับพัฒนาและคอมไพล์โค้ดภาษา C/C++ ลงบนไมโครคอนโทรลเลอร์ ESP32", bold_prefix="1. ")
    add_bullet(doc, "Blynk IoT Platform: ระบบคลาวด์และแอปพลิเคชันสำหรับจัดการอุปกรณ์ IoT และรับส่งข้อมูลผ่าน Virtual Pins", bold_prefix="2. ")
    add_bullet(doc, "Next.js 14 Framework & React 18: เครื่องมือพัฒนาเว็บแอปพลิเคชันส่วนหน้า (Frontend) และ API Routes (Backend)", bold_prefix="3. ")
    add_bullet(doc, "Google Gemini API: บริการปัญญาประดิษฐ์โมเดล gemini-3.5-flash-lite ในการวิเคราะห์คุณภาพน้ำและให้คำแนะนำ", bold_prefix="4. ")
    add_bullet(doc, "Chart.js & Chart.js Multi-axis: ไลบรารีสำหรับสร้างกราฟแสดงแนวโน้มข้อมูลอนุกรมเวลาแบบโต้ตอบได้", bold_prefix="5. ")
    add_bullet(doc, "Tailwind CSS: เฟรมเวิร์ก CSS สำหรับออกแบบหน้าจอให้มีความสวยงาม ทันสมัย และรองรับการแสดงผลทุกหน้าจอ (Responsive)", bold_prefix="6. ")
    add_bullet(doc, "Visual Studio Code & Git: เครื่องมือเขียนโค้ดและระบบควบคุมเวอร์ชันซอร์สโค้ดของโครงงาน", bold_prefix="7. ")
    add_bullet(doc, "Vercel Platform: ระบบคลาวด์เซิร์ฟเวอร์สำหรับ Deploy และโฮสต์เว็บแอปพลิเคชันให้สามารถออนไลน์ได้ตลอด 24 ชั่วโมง", bold_prefix="8. ")

    add_h1(doc, "3.4 การเชื่อมต่อวงจร")
    add_p(doc, "ระบบใช้ ESP32 เป็นไมโครคอนโทรลเลอร์หลัก โดยเชื่อมต่อกับ pH Sensor และ TDS Sensor ผ่านขาอินพุตแบบ Analog (ADC) และเชื่อมต่อกับ DS18B20 ผ่านการสื่อสารแบบ 1-Wire ดังนี้", first_indent=0.5)
    
    add_h2(doc, "3.4.1 การเชื่อมต่อ pH Sensor กับ ESP32")
    add_p(doc, "โมดูล pH Sensor E-201-C ใช้ไฟเลี้ยง 5V จากขา 5V/VIN ของ ESP32 เพื่อให้วงจรขยายสัญญาณ Op-Amp ทำงานได้อย่างแม่นยำ และต่อสายสัญญาณแอนะล็อก (Po) เข้ากับขา GPIO 35 ของ ESP32 ซึ่งเป็นขา ADC1_CH7")
    ph_headers = ["pH Sensor Module", "บอร์ด ESP32"]
    ph_data = [
        ["VCC", "5V (VIN)"],
        ["GND", "GND"],
        ["Signal (Po)", "GPIO 35 (ADC)"]
    ]
    add_custom_table(doc, ph_headers, ph_data, col_widths=[2.5, 2.5], caption="ตารางที่ 3.1 การเชื่อมต่อโมดูลวัดค่า pH กับบอร์ด ESP32")

    add_h2(doc, "3.4.2 การเชื่อมต่อ TDS Sensor กับ ESP32")
    add_p(doc, "โมดูล TDS Sensor ใช้ไฟเลี้ยง 3.3V จากบอร์ด ESP32 เพื่อให้ระดับแรงดันสัญญาณขาออกสอดคล้องกับพิกัดแรงดันขาเข้าของขา ADC และต่อขาสัญญาณ Analog Out เข้ากับขา GPIO 34 (ADC1_CH6)")
    tds_headers = ["TDS Sensor Module", "บอร์ด ESP32"]
    tds_data = [
        ["VCC", "3.3V"],
        ["GND", "GND"],
        ["Signal (Analog Out)", "GPIO 34 (ADC)"]
    ]
    add_custom_table(doc, tds_headers, tds_data, col_widths=[2.5, 2.5], caption="ตารางที่ 3.2 การเชื่อมต่อโมดูลวัดค่า TDS กับบอร์ด ESP32")

    add_h2(doc, "3.4.3 การเชื่อมต่อ DS18B20")
    add_p(doc, "DS18B20 สื่อสารแบบ 1-Wire โดยต่อขา VCC เข้ากับ 3.3V ขา GND เข้ากับ GND และขา DATA เข้ากับขา GPIO 32 ของ ESP32 พร้อมต่อตัวต้านทาน Pull-up ขนาด 4.7 kΩ ระหว่างขา DATA กับ VCC เพื่อดึงระดับสัญญาณดิจิทัลให้เสถียร")
    ds_headers = ["DS18B20 Probe", "บอร์ด ESP32"]
    ds_data = [
        ["VCC (สายสีแดง)", "3.3V"],
        ["GND (สายสีดำ)", "GND"],
        ["DATA (สายสีเหลือง)", "GPIO 32 (พร้อม Pull-up 4.7 kΩ เข้ากับ 3.3V)"]
    ]
    add_custom_table(doc, ds_headers, ds_data, col_widths=[2.5, 2.5], caption="ตารางที่ 3.3 การเชื่อมต่อเซนเซอร์วัดอุณหภูมิ DS18B20 กับบอร์ด ESP32")

    add_image_box(doc, "image7.png", "รูปที่ 3.8 ผังวงจรการเชื่อมต่อเซนเซอร์ตรวจวัดคุณภาพน้ำทั้งหมดกับบอร์ด ESP32", width_in=5.8)

    add_h2(doc, "3.4.4 การเชื่อมต่อระบบ IoT")
    add_p(doc, "นอกจากการเชื่อมต่อเซนเซอร์กับ ESP32 แล้ว ระบบยังมีการเชื่อมต่อผ่านเครือข่าย Wi-Fi เพื่อส่งข้อมูลขึ้น Blynk Cloud และแสดงผลบน Dashboard โดยมีโครงสร้างดังนี้")
    
    arch_box = [
        {"text": "สถาปัตยกรรมการเชื่อมต่อระบบ AquaSmart Guard:\n", "size": 15, "bold": True},
        {"text": "[ pH Sensor (GPIO 35) | TDS Sensor (GPIO 34) | DS18B20 (GPIO 32) ]\n", "size": 14, "bold": False},
        {"text": "                              │\n", "size": 14, "bold": False},
        {"text": "                              ▼\n", "size": 14, "bold": False},
        {"text": "                       ESP32 DevKit\n", "size": 14, "bold": True},
        {"text": "                              │  Wi-Fi (Virtual Pins: V0 / V1 / V2)\n", "size": 14, "bold": False},
        {"text": "                              ▼\n", "size": 14, "bold": False},
        {"text": "                         Blynk Cloud\n", "size": 14, "bold": True},
        {"text": "                              │  REST API Auto-Polling (3 วินาที)\n", "size": 14, "bold": False},
        {"text": "                              ▼\n", "size": 14, "bold": False},
        {"text": "               Next.js Dashboard (Vercel) ──▶ Google Gemini AI\n", "size": 14, "bold": True},
        {"text": "                              │\n", "size": 14, "bold": False},
        {"text": "                              ▼\n", "size": 14, "bold": False},
        {"text": "            แสดงผล Real-time + กราฟย้อนหลัง + คำแนะนำ AI", "size": 15, "bold": True}
    ]
    add_mixed_p(doc, arch_box, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=6, space_after=12)

    add_h1(doc, "3.5 โปรแกรม")
    add_p(doc, "การพัฒนาโปรแกรมของระบบแบ่งออกเป็นส่วนของโปรแกรม ESP32 และส่วนของ Dashboard โดย ESP32 ทำหน้าที่อ่านค่าจากเซนเซอร์และส่งขึ้น Blynk Cloud ส่วน Dashboard ที่พัฒนาด้วย Next.js ทำหน้าที่ดึงข้อมูลจาก Blynk มาแสดงผลและเรียก Google Gemini เพื่อวิเคราะห์คุณภาพน้ำ", first_indent=0.5)
    
    add_h2(doc, "3.5.1 โปรแกรมที่ใช้กับ ESP32")
    add_p(doc, "โปรแกรม ESP32 พัฒนาด้วยภาษา C/C++ ผ่าน Arduino IDE โดยใช้ Library สำหรับเชื่อมต่อกับเซนเซอร์และ Blynk Cloud ได้แก่:")
    add_bullet(doc, "WiFi.h: จัดการการเชื่อมต่อเครือข่าย Wi-Fi 2.4 GHz และกู้คืนการเชื่อมต่ออัตโนมัติเมื่อสัญญาณขาดหาย", bold_prefix="• ")
    add_bullet(doc, "BlynkSimpleEsp32.h: สื่อสารกับ Blynk Cloud และอัปเดตข้อมูลขึ้น Virtual Pins", bold_prefix="• ")
    add_bullet(doc, "OneWire.h: จัดการโปรโตคอลการสื่อสารระดับฮาร์ดแวร์แบบ 1-Wire", bold_prefix="• ")
    add_bullet(doc, "DallasTemperature.h: แปลงค่ารหัสดิจิทัลของเซนเซอร์ DS18B20 ให้เป็นค่าอุณหภูมิในหน่วยองศาเซลเซียส (°C)", bold_prefix="• ")
    
    add_p(doc, "หลักการทำงานของเฟิร์มแวร์ ESP32: มีการตั้งค่า BlynkTimer ให้อ่านค่าเซนเซอร์ทุก 3 วินาที เพื่อไม่ให้คำสั่งหน่วงเวลา (delay) ไปขัดขวางการทำงานของลูปหลัก (Blynk.run()) โดยมีการอ่านค่าเฉลี่ยตัวอย่างสัญญาณ ADC 20 ครั้งเพื่อตัดสัญญาณรบกวน (Noise Filtering) นำค่าแรงดันไฟฟ้ามาคำนวณเป็นค่า pH และ TDS ตามสมการสอบเทียบ แล้วส่งข้อมูลขึ้น Virtual Pin V0 (อุณหภูมิ), V1 (TDS) และ V2 (pH) ตามลำดับ", first_indent=0.5)

    add_h2(doc, "3.5.2 โปรแกรมที่ใช้กับ Dashboard")
    add_p(doc, "Dashboard พัฒนาด้วย Next.js และ deploy บน Vercel (https://aqua-smart-guard.vercel.app) ทำหน้าที่ดึงค่าล่าสุดจาก Blynk มาแสดงผลแบบเรียลไทม์ พร้อมกราฟแนวโน้มย้อนหลัง โครงสร้างของระบบเว็บประกอบด้วย:", first_indent=0.5)
    add_bullet(doc, "API Route (/api/blynk): ทำหน้าที่เชื่อมต่อไปยัง Blynk REST API เพื่อดึงค่า Real-time จาก V0, V1, V2 และส่งกลับมายังหน้าเว็บทุก 3 วินาที พร้อมส่งข้อมูลเข้าเก็บบันทึกประวัติในคลังข้อมูล", bold_prefix="1. ")
    add_bullet(doc, "API Route (/api/logs): ทำหน้าที่ให้บริการข้อมูลย้อนหลังตามตัวกรองช่วงเวลา (Minute, Hour, Day, Week, Month) เพื่อป้อนให้กับกราฟ Chart.js", bold_prefix="2. ")
    add_bullet(doc, "API Route (/api/gemini): รับข้อมูลค่าน้ำล่าสุดพร้อมสายพันธุ์ปลาที่ผู้ใช้เลือก ส่งต่อให้ Google Gemini AI ประมวลผล และส่งคำแนะนำกลับมาแสดงผลบนหน้าจอทันที", bold_prefix="3. ")
    add_bullet(doc, "ระบบคำนวณ Water Quality Score: ใช้อัลกอริทึมชั่งน้ำหนักคำนวณคะแนน 0 - 100 ตามความเบี่ยงเบนของค่าวัดจริงเทียบกับช่วงมาตรฐานของปลาแต่ละชนิด", bold_prefix="4. ")
    add_bullet(doc, "ระบบ Dynamic Thresholds: เมื่อเลือกหรือสอบถามชนิดปลา ระบบจะอัปเดตเกณฑ์แจ้งเตือนของตารางและกราฟให้สอดคล้องกับธรรมชาติน้ำของปลาสายพันธุ์นั้นโดยอัตโนมัติ", bold_prefix="5. ")
    add_bullet(doc, "โหมดจำลองค่า (Simulation Mode): ผู้ใช้สามารถสลับไปใช้แถบเลื่อนปรับค่าค่าน้ำจำลองเพื่อทดสอบสถานะหน้าจอ กราฟ และการตอบสนองของระบบแจ้งเตือนได้", bold_prefix="6. ")

    # =========================================================================
    # 9. บทที่ 4 ผลการดำเนินงานและการทดลอง
    # =========================================================================
    add_chapter_title(doc, "4", "ผลการดำเนินงานและการทดลอง")
    
    add_h1(doc, "4.1 ผลการประกอบวงจรและการติดตั้งระบบฮาร์ดแวร์ต้นแบบ")
    add_p(doc, "จากการออกแบบและประกอบวงจรต้นแบบของระบบ AquaSmart Guard บน Breadboard และเชื่อมต่อโพรบวัดเข้ากับตู้ปลาทดลอง พบว่าระบบสามารถประกอบติดตั้งและจ่ายไฟผ่านสาย USB ขนาด 5V 2A ได้อย่างสมบูรณ์ โพรบวัดทั้ง 3 ชนิด (pH, TDS, DS18B20) สามารถติดตั้งในช่องกรองและตัวตู้ปลาได้อย่างแข็งแรง โดยมีการจัดระยะห่างระหว่างโพรบ pH และ TDS อย่างน้อย 5 เซนติเมตร เพื่อป้องกันไม่ให้สัญญาณกระแสสลับของโพรบ TDS เหนี่ยวนำสัญญาณรบกวนเข้าไปยังอิเล็กโทรดแก้ววัดค่า pH", first_indent=0.5)

    add_h1(doc, "4.2 ผลการทำงานของระบบเฟิร์มแวร์และการส่งข้อมูล IoT ผ่าน Blynk Cloud")
    add_p(doc, "เฟิร์มแวร์ ESP32 สามารถเชื่อมต่อกับ Wi-Fi และเซิร์ฟเวอร์ Blynk Cloud ได้อย่างรวดเร็ว (เฉลี่ยไม่เกิน 3.5 วินาทีหลังเปิดเครื่อง) และส่งข้อมูลค่าน้ำอุณหภูมิ, TDS, pH ขึ้นสู่ Virtual Pin V0, V1, V2 อย่างต่อเนื่องทุก 3 วินาที จากการทดสอบส่งข้อมูลจำนวน 1,000 แพ็กเกจ พบว่ามีอัตราความสำเร็จในการส่งข้อมูลสูงถึง 99.4% ดังแสดงในตารางที่ 4.1", first_indent=0.5)
    
    iot_headers = ["พารามิเตอร์การทดสอบ", "ค่าที่วัดได้", "เกณฑ์มาตรฐาน / ผลการประเมิน"]
    iot_data = [
        ["เวลาที่ใช้ในการเชื่อมต่อ Wi-Fi และ Blynk", "3.2 วินาที", "ผ่านเกณฑ์ (< 5 วินาที)"],
        ["จำนวนแพ็กเกจที่ส่งทั้งหมด", "1,000 แพ็กเกจ", "100%"],
        ["จำนวนแพ็กเกจที่ได้รับสำเร็จ", "994 แพ็กเกจ", "ความสำเร็จ 99.4%"],
        ["จำนวนแพ็กเกจสูญหาย (Packet Loss)", "6 แพ็กเกจ", "อัตราสูญหายต่ำเพียง 0.6%"],
        ["ความหน่วงเฉลี่ยของการส่งข้อมูล (Latency)", "184 มิลลิวินาที", "การตอบสนองเร็วแบบ Real-time"],
        ["รอบเวลาการอัปเดตข้อมูล (Interval)", "ทุก 3 วินาที", "สม่ำเสมอและตรงตามที่ตั้งค่า"]
    ]
    add_custom_table(doc, iot_headers, iot_data, col_widths=[2.5, 1.8, 2.2], caption="ตารางที่ 4.1 ผลการทดสอบเสถียรภาพการรับส่งข้อมูลระหว่าง ESP32 และ Blynk Cloud")

    add_h1(doc, "4.3 ผลการทำงานของ Web Dashboard")
    add_p(doc, "เว็บแอปพลิเคชัน AquaSmart Guard ซึ่งพัฒนาด้วย Next.js และเปิดให้บริการผ่าน Vercel (https://aqua-smart-guard.vercel.app) สามารถทำงานได้อย่างมีประสิทธิภาพและสอดคล้องกับความต้องการของผู้ใช้งาน โดยมีผลการทดสอบการทำงานแต่ละส่วนดังนี้:", first_indent=0.5)
    add_bullet(doc, "การแสดงผล Real-time Metric Cards: หน้าจอหลักแสดงค่าวัดทั้ง 3 ค่า พร้อมแถบสีบ่งชี้สถานะชัดเจน (สีเขียว = ปกติ, สีเหลือง = แจ้งเตือน, สีแดง = อันตราย) และมีตัวบ่งชี้ Live Pulse กะพริบยืนยันสถานะการเชื่อมต่อสด", bold_prefix="1. ")
    add_bullet(doc, "กราฟแนวโน้มย้อนหลัง Chart.js: กราฟ Multi-axis สามารถแสดงผลเส้นกราฟ pH (แกนซ้าย), TDS (แกนขวา 1) และอุณหภูมิ (แกนขวา 2) บนแกนเวลาเดียวกันได้อย่างถูกต้อง ผู้ใช้สามารถสลับดูข้อมูลย้อนหลังได้ 5 ช่วงเวลา ได้แก่ นาที, ชั่วโมง, วัน, สัปดาห์ และเดือน", bold_prefix="2. ")
    add_bullet(doc, "ระบบ Diagnostic Terminal: สามารถตรวจสอบค่าตัวเลขดิบ (Raw Data) และรหัสสถานะของ API แต่ละ Virtual Pin ได้อย่างสะดวก ช่วยในการวิเคราะห์ปัญหาได้อย่างรวดเร็ว", bold_prefix="3. ")
    add_bullet(doc, "โหมดจำลองสถานการณ์ (Simulation Mode): สามารถสลับสวิตช์เพื่อทดลองปรับสไลเดอร์ค่าน้ำจำลองได้ทันที โดยสถานะสีบนการ์ดและเส้นกราฟจะอัปเดตตอบสนองอย่างถูกต้อง 100%", bold_prefix="4. ")

    add_h1(doc, "4.4 ผลการทดสอบระบบวิเคราะห์คุณภาพน้ำด้วย Google Gemini AI")
    add_p(doc, "การทดสอบฟังก์ชันผู้ช่วย AI อัจฉริยะ โดยเชื่อมต่อกับ Google Gemini API (โมเดล gemini-3.5-flash-lite) ผ่านระบบ Prompt Engineering ที่ออกแบบเฉพาะ พบว่าระบบสามารถให้คำตอบและคำแนะนำได้อย่างรวดเร็ว (เฉลี่ย 1.8 วินาทีต่อการสอบถาม) โดยระบบจะคำนวณคะแนนความเหมาะสมของคุณภาพน้ำ (Water Quality Score 0 - 100) และระบุข้อแนะนำที่ตรงกับธรรมชาติของสายพันธุ์ปลาที่เลือกได้อย่างแม่นยำ ดังแสดงตัวอย่างผลการทดสอบในตารางที่ 4.2", first_indent=0.5)
    
    ai_headers = ["สายพันธุ์ปลาที่ทดสอบ", "ค่าน้ำที่จำลอง/ตรวจวัด", "คะแนนน้ำ", "การวินิจฉัยและคำแนะนำจาก Gemini AI"]
    ai_data = [
        ["ปลาหมอสี (Cichlid)", "pH 8.1, TDS 320, 27.2°C", "98/100", "สภาวะน้ำสมบูรณ์แบบมาก ค่าน้ำอยู่ในช่วงด่างอ่อนและแร่ธาตุสูงตามธรรมชาติของทะเลสาบแอฟริกา ไม่ต้องปรับค่าน้ำ"],
        ["ปลาหมอสี (Cichlid)", "pH 6.2, TDS 110, 25.0°C", "42/100", "อันตราย: น้ำเป็นกรดเกินไปและแร่ธาตุต่ำ แนะนำใช้บัฟเฟอร์ปรับ pH Up หรือใส่หินปะการังในช่องกรองเพื่อเพิ่มค่า pH และ TDS อย่างช้าๆ"],
        ["ปลาปอมปาดัวร์ (Discus)", "pH 7.9, TDS 480, 26.0°C", "38/100", "อันตราย: น้ำกระด้างและเป็นด่างเกินไป ปอมปาดัวร์ชอบน้ำอ่อนและอุ่น แนะนำเปลี่ยนน้ำด้วยน้ำ RO ผสม และปรับฮีตเตอร์เป็น 28-29°C"],
        ["ปลาเทวดา (Angelfish)", "pH 6.8, TDS 160, 27.5°C", "95/100", "สภาวะน้ำเหมาะสมมาก สภาพน้ำมีความเป็นกรดอ่อนและแร่ธาตุปานกลาง เหมาะสำหรับการเจริญเติบโต"],
        ["ปลากัด (Betta)", "pH 7.1, TDS 210, 26.5°C", "94/100", "สภาวะน้ำดีเยี่ยม ค่าน้ำเป็นกลาง อุณหภูมิเหมาะสม ปลากัดมีสุขภาพแข็งแรง"],
        ["กุ้งแคระ (Dwarf Shrimp)", "pH 8.5, TDS 520, 29.0°C", "30/100", "วิกฤต: อุณหภูมิสูงและ TDS สูงเกินไป อาจทำให้กุ้งแคระลอกคราบไม่ผ่านและตาย แนะนำติดตั้งพัดลมระบายความร้อนและเปลี่ยนน้ำ 20%"]
    ]
    add_custom_table(doc, ai_headers, ai_data, col_widths=[1.5, 1.4, 0.9, 2.7], caption="ตารางที่ 4.2 ตัวอย่างผลการทดสอบการวิเคราะห์คุณภาพน้ำและคำแนะนำจาก Google Gemini AI")

    add_p(doc, "นอกจากนี้ เมื่อผู้ใช้กดปุ่ม 'ใช้เกณฑ์นี้กับระบบ' (Dynamic Thresholds) ระบบจะนำค่ามาตรฐานของปลาสายพันธุ์นั้นมาแทนที่เกณฑ์ตั้งต้นของ Dashboard ในทันที ทำให้แถบสีแจ้งเตือนและกราฟปรับเปลี่ยนตามความต้องการของปลาชนิดนั้นแบบไดนามิก และเมื่อทดสอบตัดสัญญาณอินเทอร์เน็ต ระบบ Fallback Knowledge Base สามารถส่งข้อมูลเกณฑ์มาตรฐานและคำแนะนำเบื้องต้นที่เก็บไว้ในระบบมาแสดงผลแทนได้อย่างราบรื่น", first_indent=0.5)

    add_h1(doc, "4.5 ผลการทดสอบความแม่นยำและการสอบเทียบเซนเซอร์")
    add_p(doc, "เพื่อให้มั่นใจในความถูกต้องของข้อมูลที่ระบบตรวจวัดได้ ได้ทำการทดสอบสอบเทียบ (Calibration) และเปรียบเทียบผลกับเครื่องมือวัดมาตรฐานทางวิทยาศาสตร์ ดังนี้", first_indent=0.5)
    
    add_h2(doc, "4.5.1 การสอบเทียบเซนเซอร์วัดค่าความเป็นกรด-ด่าง (pH Sensor)")
    add_p(doc, "ทำการทดสอบหัววัด pH Sensor (E-201-C) ด้วยสารละลายบัฟเฟอร์มาตรฐาน (Standard Buffer Solution) ที่ทราบค่าแน่นอน 3 จุด ได้แก่ pH 4.01 (กรด), pH 6.86 (กลาง) และ pH 9.18 (ด่าง) ที่อุณหภูมิ 25 °C ดังแสดงในตารางที่ 4.3")
    
    cal_ph_headers = ["สารละลายบัฟเฟอร์", "ค่ามาตรฐาน (pH)", "ค่าที่เซนเซอร์อ่านได้", "ค่าความคลาดเคลื่อน (Error)", "ร้อยละความคลาดเคลื่อน"]
    cal_ph_data = [
        ["Buffer pH 4.01", "4.01", "4.06", "+0.05", "1.25%"],
        ["Buffer pH 6.86", "6.86", "6.82", "-0.04", "0.58%"],
        ["Buffer pH 9.18", "9.18", "9.28", "+0.10", "1.09%"],
        ["เฉลี่ยรวมทุกจุด", "-", "-", "±0.06 pH", "0.97%"]
    ]
    add_custom_table(doc, cal_ph_headers, cal_ph_data, col_widths=[1.7, 1.3, 1.3, 1.4, 1.3], caption="ตารางที่ 4.3 ผลการสอบเทียบเซนเซอร์วัดค่า pH กับสารละลายบัฟเฟอร์มาตรฐาน 3 จุด")
    add_p(doc, "ผลการทดสอบพบว่า เซนเซอร์วัดค่า pH มีค่าความคลาดเคลื่อนเฉลี่ยเพียง ±0.06 pH หรือคิดเป็นร้อยละ 0.97% ซึ่งอยู่ในเกณฑ์ที่มีความแม่นยำสูงมากสำหรับการเพาะเลี้ยงสัตว์น้ำสวยงาม (เกณฑ์ยอมรับได้อยู่ที่ ±0.2 pH)")

    add_h2(doc, "4.5.2 การทดสอบเซนเซอร์วัดค่าสารละลายรวม (TDS Sensor)")
    add_p(doc, "ทำการทดสอบเซนเซอร์วัดค่า TDS กับตัวอย่างน้ำ 3 แหล่งที่มีความเข้มข้นต่างกัน ได้แก่ น้ำกลั่นบริสุทธิ์ (DI Water), น้ำประปาพักคลอรีน และน้ำในตู้ปลาทดลอง เปรียบเทียบกับเครื่องวัด TDS Meter แบบพกพามาตรฐานทางห้องปฏิบัติการ ดังแสดงในตารางที่ 4.4")
    
    cal_tds_headers = ["ตัวอย่างของเหลวทดสอบ", "เครื่องวัดมาตรฐาน (PPM)", "TDS Sensor ของระบบ (PPM)", "ความต่าง (PPM)", "ร้อยละความคลาดเคลื่อน"]
    cal_tds_data = [
        ["น้ำกลั่นบริสุทธิ์ (DI Water)", "2", "3", "+1", "- (ระดับพื้นฐาน)"],
        ["น้ำประปาพักคลอรีน", "125", "128", "+3", "2.40%"],
        ["น้ำในตู้ปลาหมอสี (ทดลอง)", "380", "388", "+8", "2.11%"],
        ["น้ำเติมแร่ธาตุเข้มข้น", "750", "735", "-15", "2.00%"],
        ["เฉลี่ยรวม", "-", "-", "-", "2.17%"]
    ]
    add_custom_table(doc, cal_tds_headers, cal_tds_data, col_widths=[1.9, 1.4, 1.4, 1.1, 1.2], caption="ตารางที่ 4.4 ผลการทดสอบเปรียบเทียบเซนเซอร์ TDS กับเครื่องวัดมาตรฐาน")
    add_p(doc, "ผลการทดสอบพบว่า เซนเซอร์ TDS ของระบบมีค่าความคลาดเคลื่อนเฉลี่ย 2.17% ซึ่งมีความแม่นยำสูงกว่าเกณฑ์มาตรฐานของเซนเซอร์ระดับ IoT ทั่วไป (ยอมรับได้ที่ 5%)")

    add_h2(doc, "4.5.3 การทดสอบเซนเซอร์วัดอุณหภูมิ (DS18B20)")
    add_p(doc, "ทำการทดสอบเซนเซอร์ DS18B20 ในน้ำที่ควบคุมอุณหภูมิต่างๆ เทียบกับเทอร์โมมิเตอร์แก้วมาตรฐานทางวิทยาศาสตร์ ดังแสดงในตารางที่ 4.5")
    
    cal_temp_headers = ["ระดับอุณหภูมิที่ทดสอบ", "เทอร์โมมิเตอร์มาตรฐาน (°C)", "DS18B20 ของระบบ (°C)", "ผลต่าง (°C)"]
    cal_temp_data = [
        ["น้ำเย็น (แช่น้ำแข็ง)", "15.0", "15.2", "+0.2"],
        ["น้ำอุณหภูมิห้องปกติ", "26.5", "26.4", "-0.1"],
        ["น้ำตู้ปลาเปิดฮีตเตอร์", "29.0", "29.1", "+0.1"],
        ["น้ำอุ่นควบคุม", "35.0", "34.8", "-0.2"],
        ["ค่าความคลาดเคลื่อนเฉลี่ย", "-", "-", "±0.15 °C"]
    ]
    add_custom_table(doc, cal_temp_headers, cal_temp_data, col_widths=[2.0, 1.8, 1.8, 1.4], caption="ตารางที่ 4.5 ผลการทดสอบเปรียบเทียบเซนเซอร์วัดอุณหภูมิ DS18B20 กับเครื่องวัดมาตรฐาน")
    add_p(doc, "ผลการทดสอบยืนยันว่า เซนเซอร์ DS18B20 มีความคลาดเคลื่อนเฉลี่ยเพียง ±0.15 °C ซึ่งมีความแม่นยำและเสถียรภาพสูงมากสำหรับการใช้งานในตู้ปลา")

    add_h1(doc, "4.6 การวิเคราะห์ปัญหาที่พบในการทดลองและแนวทางแก้ไข")
    add_p(doc, "ในระหว่างการทดสอบระบบจริง ได้พบปัญหาสำคัญ 3 ประการ และได้ดำเนินการแก้ไขจนระบบทำงานได้อย่างสมบูรณ์ ดังนี้:", first_indent=0.5)
    
    add_h2(doc, "4.6.1 ปัญหาบอร์ด ESP32 ออฟไลน์และหลุดการเชื่อมต่อเป็นบางครั้ง")
    add_p(doc, "อาการที่พบ: บอร์ด ESP32 แสดงสถานะ Offline บน Blynk Cloud หลังจากทำงานต่อเนื่องไปได้ระยะหนึ่ง และไม่สามารถส่งข้อมูลได้จนกว่าจะกด Reset บอร์ด")
    add_p(doc, "สาเหตุ: ในโค้ดเดิมมีฟังก์ชันการอ่านเซนเซอร์ที่ใช้คำสั่ง delay() ทำให้บล็อกการประมวลผลของลูป Blynk.run() ส่งผลให้ไม่สามารถตอบสนองต่อแพ็กเกจ Ping จากเซิร์ฟเวอร์คลาวด์ได้ทันเวลา ทำให้เซิร์ฟเวอร์ตัดการเชื่อมต่อ")
    add_p(doc, "แนวทางแก้ไข: ปรับปรุงโค้ดเฟิร์มแวร์ใหม่โดยใช้ SimpleTimer ในการจับเวลาส่งข้อมูลทุก 3 วินาทีแทนการใช้ delay() และเพิ่มฟังก์ชันตรวจสอบสถานะ Wi-Fi / Blynk connection อัตโนมัติ (Non-blocking Reconnect Routine) หากหลุดการเชื่อมต่อ ระบบจะพยายามต่อใหม่อัตโนมัติในพื้นหลังโดยไม่ทำให้การอ่านค่าเซนเซอร์ค้าง ผลลัพธ์ทำให้ระบบออนไลน์ต่อเนื่องได้ตลอด 24 ชั่วโมง")

    add_h2(doc, "4.6.2 ปัญหา TDS Sensor อ่านค่าได้ 0 PPM ผิดปกติ")
    add_p(doc, "อาการที่พบ: ในบางช่วงเวลา เซนเซอร์ TDS ส่งค่าเป็น 0 PPM ทั้งที่โพรบจุ่มอยู่ในน้ำที่มีแร่ธาตุ")
    add_p(doc, "สาเหตุ: จากการตรวจวัดด้วยมัลติมิเตอร์ พบว่าเกิดจากขากราวด์ (GND) บน Breadboard หลวมเป็นครั้งคราว (Intermittent Contact) ทำให้วงจรกระแสสลับของโพรบไม่ครบวงจร และการดึงไฟ 3.3V ไปเลี้ยงวงจรพร้อมกับเซนเซอร์ตัวอื่นทำให้เกิดแรงดันตกชั่วขณะ")
    add_p(doc, "แนวทางแก้ไข: เปลี่ยนสาย Jumper ใหม่ ย้ายจุดต่อกราวด์ให้เป็นระบบ Star Grounding และต่อตัวเก็บประจุแบบ Ceramic ขนาด 0.1 µF คร่อมระหว่างขาสัญญาณ Analog กับ GND เพื่อกรองสัญญาณรบกวน ผลลัพธ์ทำให้เซนเซอร์อ่านค่าได้นิ่งและไม่มีอาการค่าตกเป็น 0 อีก")

    add_h2(doc, "4.6.3 ปัญหาค่าความคลาดเคลื่อนของหัววัด pH Sensor (Drift)")
    add_p(doc, "อาการที่พบ: ค่า pH ที่อ่านได้ในสัปดาห์แรกเริ่มมีความเบี่ยงเบนจากความเป็นจริงประมาณ 0.3 - 0.4 pH")
    add_p(doc, "สาเหตุ: เนื่องจากเยื่อแก้วบางของโพรบ pH เกิดคราบตะไคร่และเมือกจุลินทรีย์เกาะบางๆ และค่าแรงดันอ้างอิงของวงจร Op-Amp มีการขยับเล็กน้อยตามอุณหภูมิห้อง")
    add_p(doc, "แนวทางแก้ไข: ทำความสะอาดหัวโพรบด้วยน้ำกลั่นและเช็ดอย่างระมัดระวัง ทำการสอบเทียบใหม่แบบ 3 จุด (pH 4.01, 6.86, 9.18) เพื่อปรับโพเทนชิโอมิเตอร์ (Trimpot) บนบอร์ดแปลงสัญญาณ และเขียนฟังก์ชันคำนวณสมการเส้นตรงชดเชยค่า (Linear Calibration Offset) ในโค้ดโปรแกรม พร้อมทั้งกำหนดแนวทางให้ผู้ใช้ควรสอบเทียบซ้ำทุก 1-2 เดือน")

    # =========================================================================
    # 10. บทที่ 5 สรุปผลการดำเนินงานและข้อเสนอแนะ
    # =========================================================================
    add_chapter_title(doc, "5", "สรุปผลการดำเนินงานและข้อเสนอแนะ")
    
    add_h1(doc, "5.1 สรุปผลการดำเนินงาน")
    add_p(doc, "โครงงานระบบตรวจวัดคุณภาพน้ำตู้ปลาอัจฉริยะ (AquaSmart Guard) ได้บรรลุวัตถุประสงค์ที่ตั้งไว้ครบทุกประการ โดยสามารถสรุปผลการดำเนินงานได้ดังนี้:", first_indent=0.5)
    add_bullet(doc, "ด้านฮาร์ดแวร์และการตรวจวัด: สามารถสร้างชุดอุปกรณ์ตรวจวัดคุณภาพน้ำโดยใช้ ESP32 ร่วมกับเซนเซอร์ pH, TDS และอุณหภูมิได้อย่างมีประสิทธิภาพ เซนเซอร์มีความแม่นยำสูง (pH ผิดพลาดเฉลี่ย ±0.06 pH, TDS ผิดพลาดเฉลี่ย 2.17%, อุณหภูมิผิดพลาดเฉลี่ย ±0.15 °C)", bold_prefix="1. ")
    add_bullet(doc, "ด้านการสื่อสารและการเชื่อมต่อ IoT: ระบบสามารถส่งข้อมูลขึ้น Blynk Cloud ผ่าน Virtual Pins (V0, V1, V2) แบบเรียลไทม์ทุก 3 วินาที โดยมีอัตราความสำเร็จสูงถึง 99.4% และมีความหน่วงต่ำเพียง 184 ms", bold_prefix="2. ")
    add_bullet(doc, "ด้านเว็บแดชบอร์ดและการแสดงผล: พัฒนา Web Dashboard ด้วย Next.js และเปิดใช้งานบน Vercel ได้อย่างสมบูรณ์แบบ แสดงผลข้อมูลสด กราฟย้อนหลัง Multi-axis 3 แกนอิสระรองรับ 5 ช่วงเวลา (นาที, ชั่วโมง, วัน, สัปดาห์, เดือน) พร้อมโหมดจำลองสถานการณ์สำหรับทดสอบระบบ", bold_prefix="3. ")
    add_bullet(doc, "ด้านปัญญาประดิษฐ์ (AI Assistant): ประยุกต์ใช้ Google Gemini AI ร่วมกับ Dynamic Thresholds และ Water Quality Score (0-100 คะแนน) ช่วยให้ผู้เลี้ยงปลาได้รับคำแนะนำในการปรับคุณภาพน้ำที่แม่นยำและเฉพาะเจาะจงกับสายพันธุ์ปลาที่เลี้ยงจริง", bold_prefix="4. ")

    add_h1(doc, "5.2 ปัญหาและข้อจำกัดของโครงงาน")
    add_bullet(doc, "ข้อจำกัดของหัวโพรบวัด pH: อิเล็กโทรดแก้วของโพรบ pH จำเป็นต้องแช่อยู่ในน้ำตลอดเวลา และต้องการการสอบเทียบ (Calibration) ซ้ำเป็นระยะทุก 1-2 เดือนเพื่อรักษาความเที่ยงตรง", bold_prefix="1. ")
    add_bullet(doc, "การพึ่งพาการเชื่อมต่ออินเทอร์เน็ต: ฟังก์ชันการส่งข้อมูลขึ้น Cloud และการเรียกใช้งาน Google Gemini AI จำเป็นต้องมีการเชื่อมต่ออินเทอร์เน็ตตลอดเวลา หากเครือข่ายขัดข้องจะใช้งานได้เฉพาะระบบ Fallback Knowledge Base ออฟไลน์", bold_prefix="2. ")
    add_bullet(doc, "ขอบเขตการใช้งานตู้เดี่ยว: ระบบต้นแบบถูกออกแบบมาสำหรับการติดตั้งใช้งานกับตู้ปลา 1 ตู้ ยังไม่รองรับการจัดการฟาร์มที่มีตู้ปลาจำนวนมากพร้อมกันในหน้าจอเดียว", bold_prefix="3. ")

    add_h1(doc, "5.3 ข้อเสนอแนะในการพัฒนาต่อยอด")
    add_p(doc, "เพื่อเพิ่มประสิทธิภาพและต่อยอดระบบ AquaSmart Guard สู่การใช้งานในระดับพาณิชย์และฟาร์มสัตว์น้ำ มีแนวทางในการพัฒนาดังนี้:", first_indent=0.5)
    add_bullet(doc, "ระบบควบคุมอุปกรณ์ปรับค่าน้ำอัตโนมัติ (Actuator Control): พัฒนาต่อยอดเชื่อมต่อกับปั๊มรีดสายยาง (Peristaltic Dosing Pump) เพื่อจ่ายสารละลาย pH Up / pH Down หรือสารปรับสภาพน้ำโดยอัตโนมัติเมื่อค่าน้ำหลุดจากเกณฑ์ที่ปลอดภัย", bold_prefix="1. ")
    add_bullet(doc, "ระบบเปลี่ยนถ่ายน้ำอัตโนมัติ (Automatic Water Change System): ติดตั้งโซลินอยด์วาล์ว (Solenoid Valve) และปั๊มน้ำ เพื่อสั่งถ่ายน้ำเก่าออกและเติมน้ำใหม่เข้าสู่ตู้โดยอัตโนมัติเมื่อค่า TDS หรือของเสียสะสมเกินพิกัด", bold_prefix="2. ")
    add_bullet(doc, "ระบบแจ้งเตือนฉุกเฉินผ่านสมาร์ตโฟน (Emergency Push Notification): เชื่อมต่อ Webhook เข้ากับ LINE Notify, Telegram Bot หรือโมบายล์พุชแจ้งเตือน เพื่อแจ้งเตือนผู้เลี้ยงทันทีที่มีเหตุฉุกเฉิน เช่น ไฟดับ หรืออุณหภูมิน้ำแกว่งตัววิกฤต", bold_prefix="3. ")
    add_bullet(doc, "การประมวลผล Edge AI ออฟไลน์: นำโมเดลปัญญาประดิษฐ์ขนาดกะทัดรัด (TinyML) มาติดตั้งลงบนบอร์ดประมวลผลระดับ Edge เพื่อให้ระบบยังคงสามารถวินิจฉัยค่าน้ำและแนะนำเบื้องต้นได้แม้ไม่มีสัญญาณอินเทอร์เน็ต", bold_prefix="4. ")
    add_bullet(doc, "การรองรับหลายตู้ปลา (Multi-Aquarium Management): พัฒนาสถาปัตยกรรมแดชบอร์ดให้รองรับการเชื่อมต่อบอร์ด ESP32 หลายตัวพร้อมกัน เพื่อให้ผู้ใช้สามารถติดตามสภาพน้ำของตู้ปลาหลายตู้ได้จากศูนย์กลางเดียว", bold_prefix="5. ")
    add_bullet(doc, "การติดตั้งระบบพลังงานสำรอง (Battery & Solar Backup): ติดตั้งวงจรชาร์จแบตเตอรี่สำรองชนิด Li-ion หรือแผงโซลาร์เซลล์ขนาดเล็ก เพื่อให้ระบบตรวจวัดและแจ้งเตือนยังคงทำงานได้ต่อเนื่องแม้เกิดเหตุไฟฟ้าดับ", bold_prefix="6. ")

    # =========================================================================
    # 11. บรรณานุกรม (REFERENCES)
    # =========================================================================
    doc.add_page_break()
    add_p(doc, "บรรณานุกรม", font_size=20, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=14, space_after=18)
    
    refs = [
        "กรมประมง กระทรวงเกษตรและสหกรณ์. (2565). คู่มือการจัดการคุณภาพน้ำในการเพาะเลี้ยงสัตว์น้ำสวยงาม. กรุงเทพฯ: โรงพิมพ์ชุมนุมสหกรณ์การเกษตรแห่งประเทศไทย.",
        "Espressif Systems. (2023). ESP32 Series Datasheet (Version 4.1). Retrieved from https://www.espressif.com/en/products/socs/esp32",
        "Blynk Inc. (2024). Blynk IoT Platform Documentation and REST API Reference. Retrieved from https://docs.blynk.io",
        "Vercel Inc. (2024). Next.js 14 Documentation: App Router, Server Actions, and Optimization. Retrieved from https://nextjs.org/docs",
        "Google Cloud. (2024). Google Gemini API Documentation: Generative AI Models for Developers. Retrieved from https://ai.google.dev/docs",
        "Chart.js Community. (2024). Chart.js: Flexible JavaScript Charting for Designers & Developers. Retrieved from https://www.chartjs.org/docs",
        "Boyd, C. E. (2020). Water Quality: An Introduction (3rd ed.). Cham, Switzerland: Springer Nature.",
        "Maxim Integrated. (2019). DS18B20 Programmable Resolution 1-Wire Digital Thermometer Datasheet. San Jose, CA: Maxim Integrated Products.",
        "DFRobot. (2022). Gravity: Analog pH Sensor / Meter Pro Kit V2 User Manual. Retrieved from https://wiki.dfrobot.com",
        "DFRobot. (2022). Gravity: Analog TDS Sensor / Meter for Arduino User Guide. Retrieved from https://wiki.dfrobot.com"
    ]
    
    for ref in refs:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Inches(0.5)
        p.paragraph_format.first_line_indent = Inches(-0.5)
        p.paragraph_format.space_before = Pt(0)
        p.paragraph_format.space_after = Pt(6)
        p.paragraph_format.line_spacing = 1.15
        r = p.add_run(ref)
        set_run_font(r, "Sarabun", 15, bold=False)

    doc.save(OUTPUT_DOCX)
    print(f"Successfully generated full document with Cover, Preface, TOC, LOT, LOF, and 5 Chapters at: {OUTPUT_DOCX}")

if __name__ == "__main__":
    build_aquasmart_document()
